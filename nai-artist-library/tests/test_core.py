from __future__ import annotations

import unittest
import tempfile
import time
import zipfile
import sqlite3
import json
from datetime import datetime
from contextlib import closing
from functools import partial
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from PIL import Image, PngImagePlugin
from docx import Document

from app import app
from database import connect, entry_media_dir, init_db
from docx_database_converter import (
    apply_incremental_update,
    build_incremental_plan,
    export_database,
    import_documents,
    preview_import,
    read_markdown_document,
    read_scene_document,
)
from image_metadata import extract_image_metadata
import integrated_features
from library_audit import audit_library
from importer import (
    Paragraph,
    ParsedDocx,
    artist_blocks,
    artist_header,
    looks_sensitive,
    prompt_blocks,
    rating_from_text,
    split_mapping,
    split_positive_negative,
    common_negative_entries,
)


class ImportRulesTest(unittest.TestCase):
    def test_half_star_rating(self):
        self.assertEqual(rating_from_text("★★★★☆"), 9)
        self.assertEqual(rating_from_text("★★★"), 6)
        self.assertIsNone(rating_from_text("尚未评分"))

    def test_sensitive_tokens_are_detected(self):
        self.assertTrue(looks_sensitive("sk-abcdefghijklmnopqrstuvwxyz123456"))
        self.assertTrue(looks_sensitive("pst-abcdefghijklmnopqrstuvwxyz123456"))
        self.assertFalse(looks_sensitive("artist:example, best quality"))

    def test_title_prompt_mapping(self):
        self.assertEqual(split_mapping("场景名称：tag one, tag two"), ("场景名称", "tag one, tag two"))

    def test_title_prompt_mapping_accepts_word_line_breaks(self):
        self.assertEqual(split_mapping("画师名称：tag one,\ntag two"), ("画师名称", "tag one, tag two"))

    def test_title_prompt_mapping_accepts_weight_syntax_at_end(self):
        self.assertEqual(split_mapping("画师名称：3::clean lines::"), ("画师名称", "3::clean lines::"))

    def test_title_prompt_mapping_accepts_long_titles(self):
        title = "与巨大怪物（monster可以改为Mon3tr(arknights)）"
        self.assertEqual(split_mapping(f"{title}：nsfw, 1girl"), (title, "nsfw, 1girl"))

    def test_title_prompt_mapping_accepts_empty_prompts(self):
        self.assertEqual(split_mapping("有点微妙的画风："), ("有点微妙的画风", ""))

    def test_title_prompt_mapping_accepts_commas_in_titles(self):
        self.assertEqual(split_mapping("双头龙,loli构图：2girls, yuri"), ("双头龙,loli构图", "2girls, yuri"))

    def test_artist_name_uses_full_width_colon(self):
        parsed = artist_header("★★★★★瑟瑟手绘漫画：1.8::artist one::, artist:two")
        self.assertEqual(parsed, (10, "瑟瑟手绘漫画", "1.8::artist one::, artist:two"))

    def test_rated_artist_also_accepts_ascii_colon(self):
        self.assertEqual(artist_header("★★★★☆未命名7:artist one, artist two"), (9, "未命名7", "artist one, artist two"))
        self.assertEqual(artist_header("柔美厚涂混合: artist one, best quality"), (None, "柔美厚涂混合", "artist one, best quality"))
        self.assertIsNone(artist_header("1.2::artist:name, best quality"))

    def test_real_negative_marker_variants_keep_the_prompt(self):
        positive, negative, _ = split_positive_negative([
            "artist one, best quality",
            "负面(worst quality:1.2), lowres",
        ])
        self.assertEqual(positive, "artist one, best quality")
        self.assertEqual(negative, "(worst quality:1.2), lowres")

        _, negative, _ = split_positive_negative(["暂无负面，暂时用blurry, lowres"])
        self.assertEqual(negative, "blurry, lowres")

        _, negative, _ = split_positive_negative(["负面（或者通用）：bad anatomy"])
        self.assertEqual(negative, "bad anatomy")

    def test_artist_field_label_is_not_a_title(self):
        self.assertIsNone(artist_header("负面：bad anatomy, low quality"))

    def test_consecutive_named_scenes_are_separate_cards(self):
        doc = ParsedDocx(
            Path("sample.docx"),
            [
                Paragraph(0, "正戏：", outline_level=1),
                Paragraph(1, "隐秘：", outline_level=2),
                Paragraph(2, "场景甲：tag one, tag two"),
                Paragraph(3, "场景乙：tag three, tag four"),
            ],
            {},
        )
        entries = prompt_blocks(doc)
        self.assertEqual([item["title"] for item in entries], ["场景甲", "场景乙"])
        self.assertTrue(all(item["category"] == "正戏/隐秘" for item in entries))

    def test_three_images_then_three_artists_are_paired(self):
        doc = ParsedDocx(
            Path("sample.docx"),
            [
                Paragraph(0, "", ["rId1", "rId2", "rId3"]),
                Paragraph(1, "★★★★★画师甲：artist one"),
                Paragraph(2, "★★★★☆画师乙：artist two"),
                Paragraph(3, "★★★★画师丙：artist three"),
            ],
            {},
        )
        entries = artist_blocks(doc)
        self.assertEqual([len(item["image_rel_ids"]) for item in entries], [1, 1, 1])
        self.assertEqual([item["title"] for item in entries], ["画师甲", "画师乙", "画师丙"])

    def test_unrated_artist_stays_unrated_and_does_not_borrow_an_image(self):
        doc = ParsedDocx(
            Path("sample.docx"),
            [
                Paragraph(0, "", ["rId1"]),
                Paragraph(1, "无评分画师：artist one, best quality, detailed lighting"),
                Paragraph(2, "负面：bad anatomy"),
                Paragraph(3, "★★★★★下一位：artist two"),
            ],
            {},
        )
        entries = artist_blocks(doc)
        self.assertEqual([item["rating"] for item in entries], [None, 10])
        self.assertEqual(entries[0]["negative_prompt"], "bad anatomy")
        self.assertEqual([len(item["image_rel_ids"]) for item in entries], [1, 0])

    def test_explicit_standalone_negative_section_creates_three_unillustrated_cards(self):
        doc = ParsedDocx(
            Path("sample.docx"),
            [
                Paragraph(0, "", ["rId1", "rId2", "rId3"]),
                Paragraph(1, "★★★★★画师甲：artist one"),
                Paragraph(2, "以下应该被认为是三个单独的负面提示词默认负面：bad one"),
                Paragraph(3, "通用负面1（还不错）："),
                Paragraph(4, "bad two"),
                Paragraph(5, "通用负面2（较强）："),
                Paragraph(6, "bad three"),
            ],
            {},
        )
        artists = artist_blocks(doc)
        negatives = common_negative_entries(doc)
        self.assertEqual(len(artists), 1)
        self.assertEqual(len(artists[0]["image_rel_ids"]), 3)
        self.assertEqual([item["title"] for item in negatives], ["默认负面", "通用负面1（还不错）", "通用负面2（较强）"])
        self.assertEqual([item["content"] for item in negatives], ["bad one", "bad two", "bad three"])
        self.assertTrue(all(not item["image_rel_ids"] and item["rating"] is None for item in negatives))


class ApiSmokeTest(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.client = app.test_client()

    def wait_cleanup(self, response):
        cleanup = response.get_json()["cleanup"]
        if not cleanup.get("queued"):
            return cleanup
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            job = self.client.get("/api/jobs/" + cleanup["job_id"]).get_json()
            if job["status"] in ("completed", "failed"):
                self.assertEqual(job["status"], "completed", job)
                return job["result"]
            time.sleep(0.02)
        self.fail("cleanup timed out")

    def test_home_and_stats(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        response.close()
        response = self.client.get("/api/stats")
        self.assertEqual(response.status_code, 200)
        self.assertIn("entries", response.get_json())

        response = self.client.get("/api/version")
        self.assertEqual(response.status_code, 200)
        self.assertIn("version", response.get_json())

        response = self.client.get("/api/settings")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Path(response.get_json()["media_dir"]).name, "media")

        response = self.client.get("/api/navigation")
        self.assertEqual(response.status_code, 200)
        self.assertTrue({"all_count", "at_least_9", "exactly_8", "between_6_and_7", "at_most_5", "unrated"}.issubset(response.get_json()["ratings"]))

        response = self.client.get("/api/external-libraries/sources")
        self.assertEqual(response.status_code, 200)
        self.assertIsInstance(response.get_json()["sources"], list)
        response.close()

    def test_library_import_draft_is_consumed_once(self):
        payload = {
            "type": "nyanovel-library-import",
            "requestId": "draft-test-1",
            "kind": "artist",
            "content": "artist:test",
            "negativePrompt": "bad anatomy",
            "image": None,
        }
        created = self.client.post("/api/library/import-drafts", json=payload)
        self.assertEqual(created.status_code, 200)
        consumed = self.client.get("/api/library/import-drafts/draft-test-1")
        self.assertEqual(consumed.status_code, 200)
        self.assertEqual(consumed.get_json()["content"], "artist:test")
        self.assertEqual(self.client.get("/api/library/import-drafts/draft-test-1").status_code, 404)

    def test_local_gallery_fts_and_rename_keep_the_same_record(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder) / "gallery"
            root.mkdir()
            original = root / "first-picture.png"
            Image.new("RGB", (32, 32), "red").save(original)
            settings = {
                "gallery_roots": [str(root)],
                "gallery_extensions": [".png"],
                "recursive_scan": False,
                "gallery_page_size": 60,
            }
            database_path = Path(folder) / "gallery.db"
            thumbnail_path = Path(folder) / "thumbs"
            with (
                patch.object(integrated_features, "GALLERY_DB", database_path),
                patch.object(integrated_features, "GALLERY_THUMB_DIR", thumbnail_path),
                patch.object(integrated_features, "load_settings", return_value=settings),
            ):
                first = integrated_features.scan_local_gallery()
                self.assertEqual(first["added"], 1)
                # Windows keeps an image file locked while Pillow extracts metadata.
                # Drain the first scan before exercising move/rename detection.
                self.assertTrue(integrated_features.wait_for_metadata_idle(timeout=5))
                with closing(integrated_features.gallery_connection()) as conn, conn:
                    first_id = conn.execute(
                        "SELECT id FROM local_images WHERE path=?", (str(original.resolve()),)
                    ).fetchone()["id"]

                renamed = root / "renamed-picture.png"
                original.rename(renamed)
                second = integrated_features.scan_local_gallery()
                self.assertEqual(second["moved"], 1)
                self.assertEqual(second["added"], 0)
                self.assertEqual(second["removed"], 0)
                with closing(integrated_features.gallery_connection()) as conn, conn:
                    row = conn.execute("SELECT id,path FROM local_images").fetchone()
                    self.assertEqual(row["id"], first_id)
                    self.assertEqual(row["path"], str(renamed.resolve()))

                response = self.client.get("/api/local-gallery/images", query_string={"q": "renamed"})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.get_json()["total"], 1)
                response.close()
                # The production scanner intentionally parses image metadata in
                # the background. Drain earlier work before Windows removes the
                # temporary SQLite database used by this test.
                self.assertTrue(integrated_features.wait_for_metadata_idle(timeout=5))

    def test_unchanged_gallery_image_with_missing_metadata_is_repaired(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder) / "gallery"
            root.mkdir()
            image_path = root / "missing-metadata.png"
            Image.new("RGB", (48, 24), "blue").save(image_path)
            settings = {
                "gallery_roots": [str(root)],
                "gallery_extensions": [".png"],
                "recursive_scan": False,
                "gallery_page_size": 60,
            }
            database_path = Path(folder) / "gallery.db"
            thumbnail_path = Path(folder) / "thumbs"
            with (
                patch.object(integrated_features, "GALLERY_DB", database_path),
                patch.object(integrated_features, "GALLERY_THUMB_DIR", thumbnail_path),
                patch.object(integrated_features, "load_settings", return_value=settings),
            ):
                integrated_features.scan_local_gallery()
                self.assertTrue(integrated_features.wait_for_metadata_idle(timeout=5))
                with closing(integrated_features.gallery_connection()) as conn, conn:
                    conn.execute(
                        "UPDATE local_images SET width=0,height=0,prompt='',metadata_json='{}' WHERE path=?",
                        (str(image_path.resolve()),),
                    )
                repaired = integrated_features.scan_local_gallery()
                self.assertEqual(repaired["added"], 0)
                self.assertEqual(repaired["updated"], 0)
                self.assertEqual(repaired["metadata_pending"], 1)
                self.assertTrue(integrated_features.wait_for_metadata_idle(timeout=5))
                with closing(integrated_features.gallery_connection()) as conn:
                    row = conn.execute(
                        "SELECT width,height,metadata_json FROM local_images WHERE path=?",
                        (str(image_path.resolve()),),
                    ).fetchone()
                self.assertEqual((row["width"], row["height"]), (48, 24))
                self.assertIn("_scan", json.loads(row["metadata_json"]))

    def test_index_file_has_metadata_before_first_gallery_refresh(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder) / "gallery"
            root.mkdir()
            image_path = root / "deanai_123_1.png"
            info = PngImagePlugin.PngInfo()
            info.add_text("Source", "NovelAI Diffusion V4.5 Full")
            info.add_text("Comment", json.dumps({
                "prompt": "test prompt", "steps": 28, "seed": 123,
                "sampler": "k_euler_ancestral", "width": 64, "height": 96,
                "model": "nai-diffusion-4-5-full",
            }))
            Image.new("RGB", (64, 96), "purple").save(image_path, pnginfo=info)
            settings = {
                "gallery_roots": [str(root)],
                "gallery_extensions": [".png"],
                "recursive_scan": False,
                "gallery_page_size": 60,
            }
            database_path = Path(folder) / "gallery.db"
            with (
                patch.object(integrated_features, "GALLERY_DB", database_path),
                patch.object(integrated_features, "load_settings", return_value=settings),
            ):
                response = self.client.post("/api/local-gallery/index-file", json={"path": str(image_path)})
                self.assertEqual(response.status_code, 200)
                with closing(integrated_features.gallery_connection()) as conn:
                    row = conn.execute(
                        "SELECT width,height,prompt,metadata_json FROM local_images WHERE path=?",
                        (str(image_path.resolve()),),
                    ).fetchone()
                self.assertEqual((row["width"], row["height"]), (64, 96))
                self.assertEqual(row["prompt"], "test prompt")
                self.assertIn("_scan", json.loads(row["metadata_json"]))

                stats = self.client.get("/api/local-gallery/stats")
                self.assertEqual(stats.status_code, 200)
                payload = stats.get_json()
                self.assertEqual(payload["generated_total"], 1)
                self.assertEqual(len(payload["history"]), 1)
                self.assertEqual(payload["history"][0]["metadata"]["parameters"]["seed"], 123)
                self.assertEqual(payload["history"][0]["size"], image_path.stat().st_size)
                stats.close()
                response.close()

    def test_local_gallery_filters_and_marks_video_media(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder) / "gallery"
            root.mkdir()
            video = root / "sample.mp4"
            video.write_bytes(b"not-a-real-video")
            settings = {
                "gallery_roots": [str(root)],
                "gallery_extensions": [".mp4"],
                "recursive_scan": False,
                "gallery_page_size": 60,
            }
            database_path = Path(folder) / "gallery.db"
            modified_ns = int(datetime(2026, 8, 25, 12, 0).timestamp() * 1_000_000_000)
            metadata = {"parameters": {"model": "nai-test", "sampler": "k_euler", "steps": 28, "scale": 5.5}}
            with (
                patch.object(integrated_features, "GALLERY_DB", database_path),
                patch.object(integrated_features, "load_settings", return_value=settings),
            ):
                with closing(integrated_features.gallery_connection()) as conn, conn:
                    conn.execute(
                        """INSERT INTO local_images(path,name,size,modified_ns,width,height,metadata_json)
                           VALUES(?,?,?,?,?,?,?)""",
                        (str(video), video.name, video.stat().st_size, modified_ns, 1920, 1080, json.dumps(metadata)),
                    )
                response = self.client.get("/api/local-gallery/images", query_string={
                    "date_from": "2026-08-25", "date_to": "2026-08-25",
                    "model": "nai-test", "sampler": "k_euler",
                    "steps_from": "28", "steps_to": "28", "cfg_from": "5.5", "cfg_to": "5.5",
                    "orientation": "landscape", "resolution": "1920x1080",
                })
                self.assertEqual(response.status_code, 200)
                payload = response.get_json()
                self.assertEqual(payload["total"], 1)
                self.assertEqual(payload["images"][0]["media_type"], "video")
                self.assertEqual(payload["images"][0]["size"], video.stat().st_size)
                self.assertIn("nai-test", payload["models"])
                response.close()

    def test_local_gallery_page_uses_configured_size_and_preserves_offset_api(self):
        with tempfile.TemporaryDirectory() as folder:
            settings = {'gallery_roots': [], 'gallery_page_size': 60}
            with (
                patch.object(integrated_features, 'GALLERY_DB', Path(folder) / 'gallery.db'),
                patch.object(integrated_features, 'load_settings', return_value=settings),
            ):
                with closing(integrated_features.gallery_connection()) as conn, conn:
                    conn.executemany(
                        'INSERT INTO local_images(path,name,size,modified_ns,width,height,metadata_json) VALUES(?,?,?,?,?,?,?)',
                        [(str(Path(folder) / f'{i}.png'), f'{i}.png', 1, i, 512, 512, '{}') for i in range(65)],
                    )
                first = self.client.get('/api/local-gallery/images?page=1').get_json()
                second = self.client.get('/api/local-gallery/images?page=2').get_json()
                self.assertEqual(first['page_size'], 60)
                self.assertEqual(first['total'], 65)
                self.assertEqual(len(first['images']), 60)
                self.assertEqual([row['name'] for row in second['images']], ['4.png', '3.png', '2.png', '1.png', '0.png'])
                legacy = self.client.get('/api/local-gallery/images?limit=3&offset=2').get_json()
                self.assertEqual([row['name'] for row in legacy['images']], ['62.png', '61.png', '60.png'])
                settings['gallery_page_size'] = 64
                changed = self.client.get('/api/local-gallery/images?page=2').get_json()
                self.assertEqual(changed['page_size'], 64)
                self.assertEqual([row['name'] for row in changed['images']], ['0.png'])

    def test_category_prefix_keeps_kind_and_category_scope(self):
        navigation = self.client.get("/api/navigation").get_json()
        nested = next(
            (item["name"].split("/")[0] for item in navigation["categories"]["prompt"] if "/" in item["name"]),
            None,
        )
        if nested is None:
            self.skipTest("当前资料库没有二级场景目录")
        response = self.client.get("/api/entries", query_string={
            "kind": "prompt", "category_prefix": nested, "favorites_only": "0", "limit": 1000,
        })
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertGreater(payload["total"], 0)
        self.assertTrue(all(entry["kind"] == "prompt" for entry in payload["entries"]))
        self.assertTrue(all(
            entry["category"] == nested or entry["category"].startswith(f"{nested}/")
            for entry in payload["entries"]
        ))

    def test_parent_category_count_includes_all_descendants(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "category-totals.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                for title, category in (
                    ("父级直属", "NSFW"),
                    ("子级一", "NSFW/1girl"),
                    ("孙级", "NSFW/1girl/futa"),
                    ("子级二", "NSFW/单女性展示"),
                ):
                    response = self.client.post("/api/entries", json={
                        "kind": "prompt", "title": title, "category": category, "content": "test",
                    })
                    self.assertEqual(response.status_code, 201)

                navigation = self.client.get("/api/navigation").get_json()
                counts = {item["name"]: item["count"] for item in navigation["categories"]["prompt"]}
                self.assertEqual(counts["NSFW"], 4)
                self.assertEqual(counts["NSFW/1girl"], 2)
                self.assertEqual(counts["NSFW/1girl/futa"], 1)

                payload = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "category_prefix": "NSFW", "limit": 100,
                }).get_json()
                self.assertEqual(payload["total"], 4)

                self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "other-root",
                    "category": "SFW/place", "content": "test",
                })
                multi = self.client.get("/api/entries/random", query_string=[
                    ("kind", "prompt"),
                    ("category_prefix", "NSFW/1girl"),
                    ("category_prefix", "SFW"),
                ])
                self.assertEqual(multi.status_code, 200)
                self.assertIn(multi.get_json()["category"], {"NSFW/1girl", "NSFW/1girl/futa", "SFW/place"})

    def test_unrated_filter_only_returns_unrated_entries(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "unrated-filter.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                unrated = self.client.post("/api/entries", json={
                    "kind": "artist", "title": "unrated", "rating": None,
                    "category": "test", "content": "test",
                }).get_json()
                self.client.post("/api/entries", json={
                    "kind": "artist", "title": "rated", "rating": 8,
                    "category": "test", "content": "test",
                })

                response = self.client.get("/api/entries", query_string={
                    "kind": "artist", "unrated_only": "1", "limit": 1000,
                })
                self.assertEqual(response.status_code, 200)
                payload = response.get_json()
                self.assertEqual(payload["total"], 1)
                self.assertEqual([entry["id"] for entry in payload["entries"]], [unrated["id"]])

    def test_artist_rating_navigation_and_filters_use_disjoint_bands(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "rating-bands.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                for index, rating in enumerate((10, 9, 8, 7, 6, 5, None), start=1):
                    response = self.client.post("/api/entries", json={
                        "kind": "artist", "title": f"artist-{index}", "rating": rating,
                        "category": "test", "content": "test",
                    })
                    self.assertEqual(response.status_code, 201)

                ratings = self.client.get("/api/navigation").get_json()["ratings"]
                self.assertEqual(ratings["at_least_9"], 2)
                self.assertEqual(ratings["exactly_8"], 1)
                self.assertEqual(ratings["between_6_and_7"], 2)
                self.assertEqual(ratings["at_most_5"], 1)
                self.assertEqual(ratings["unrated"], 1)

                bands = (("8", "8", [8]), ("6", "7", [7, 6]), ("", "5", [5]))
                for rating_min, rating_max, expected in bands:
                    payload = self.client.get("/api/entries", query_string={
                        "kind": "artist", "rating_min": rating_min,
                        "rating_max": rating_max, "limit": 100,
                    }).get_json()
                    self.assertEqual([entry["rating"] for entry in payload["entries"]], expected)

    def test_random_entry_honors_scope_rating_and_recent_exclusion(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "random-entry.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                low = self.client.post("/api/entries", json={
                    "kind": "artist", "title": "low", "rating": 4,
                    "category": "A", "content": "low",
                }).get_json()
                high = self.client.post("/api/entries", json={
                    "kind": "artist", "title": "high", "rating": 10,
                    "category": "A", "content": "high",
                }).get_json()
                unrated = self.client.post("/api/entries", json={
                    "kind": "artist", "title": "unrated", "rating": None,
                    "category": "A", "content": "unrated",
                }).get_json()
                self.client.post("/api/entries", json={
                    "kind": "artist", "title": "other-category", "rating": 10,
                    "category": "B", "content": "other",
                })

                response = self.client.get("/api/entries/random", query_string={
                    "kind": "artist", "category": "A", "rating_min": 8,
                    "include_unrated": "1", "exclude_ids": high["id"],
                    "weighting": "rating_usage",
                })
                self.assertEqual(response.status_code, 200)
                picked = response.get_json()
                self.assertEqual(picked["id"], unrated["id"])
                self.assertNotEqual(picked["id"], low["id"])

                empty = self.client.get("/api/entries/random", query_string={
                    "kind": "prompt", "category": "missing",
                })
                self.assertEqual(empty.status_code, 404)

    def test_artist_styles_are_saved_counted_and_combine_with_rating_filters(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "artist-styles.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                records = [
                    ("paint-high", 10, "厚涂"),
                    ("paint-mid", 7, "厚涂"),
                    ("anime-mid", 6, "赛璐璐"),
                    ("unclassified", None, ""),
                ]
                for title, rating, style in records:
                    response = self.client.post("/api/entries", json={
                        "kind": "artist", "title": title, "rating": rating,
                        "style": style, "category": "test", "content": "test",
                    })
                    self.assertEqual(response.status_code, 201)

                navigation = self.client.get("/api/navigation").get_json()
                styles = {item["name"]: item for item in navigation["styles"]}
                self.assertEqual(styles["厚涂"]["all_count"], 2)
                self.assertEqual(styles["厚涂"]["at_least_9"], 1)
                self.assertEqual(styles["厚涂"]["between_6_and_7"], 1)
                self.assertEqual(styles[""]["unrated"], 1)

                payload = self.client.get("/api/entries", query_string={
                    "kind": "artist", "style": "厚涂", "rating_min": "6",
                    "rating_max": "7", "limit": 100,
                }).get_json()
                self.assertEqual([entry["title"] for entry in payload["entries"]], ["paint-mid"])

                payload = self.client.get("/api/entries", query_string={
                    "kind": "artist", "style_unclassified": "1", "limit": 100,
                }).get_json()
                self.assertEqual([entry["title"] for entry in payload["entries"]], ["unclassified"])

    def test_server_pagination_uses_limit_and_offset(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "pagination.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                for index in range(5):
                    response = self.client.post("/api/entries", json={
                        "kind": "prompt", "title": f"prompt-{index}",
                        "category": "test", "content": "test",
                    })
                    self.assertEqual(response.status_code, 201)

                first = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "sort": "title", "limit": 2, "offset": 0,
                }).get_json()
                second = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "sort": "title", "limit": 2, "offset": 2,
                }).get_json()
                self.assertEqual(first["total"], 5)
                self.assertEqual(second["total"], 5)
                self.assertEqual(len(first["entries"]), 2)
                self.assertEqual(len(second["entries"]), 2)
                self.assertTrue({entry["id"] for entry in first["entries"]}.isdisjoint(entry["id"] for entry in second["entries"]))

    def test_json_export_is_readable(self):
        response = self.client.get("/api/export/json")
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers.get("Content-Disposition", ""))
        payload = response.get_json()
        self.assertEqual(payload["format"], "nai-artist-library-export")
        self.assertIn("entries", payload)

        response = self.client.get("/api/export/artists.csv")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data.startswith(b"\xef\xbb\xbf"))
        self.assertIn("attachment", response.headers.get("Content-Disposition", ""))

    def test_mobile_export_is_portable_and_readable(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "mobile-export.db"
            init_db(test_db)
            now = "2026-08-27T00:00:00+00:00"
            with connect(test_db) as conn:
                conn.execute(
                    """
                    INSERT INTO entries
                        (kind, title, content, negative_prompt, rating, style, category, created_at, updated_at)
                    VALUES ('artist', 'mobile artist', 'artist:test', 'bad anatomy', 9,
                            'anime', 'artists/test', ?, ?)
                    """,
                    (now, now),
                )

            with patch("app.connect", partial(connect, test_db)):
                response = self.client.get("/api/export/mobile.json")

            self.assertEqual(response.status_code, 200)
            self.assertIn("attachment", response.headers.get("Content-Disposition", ""))
            payload = response.get_json()
            self.assertEqual(payload["format"], "nai-artist-library-mobile")
            self.assertEqual(payload["format_version"], 2)
            self.assertEqual(len(payload["entries"]), 1)
            entry = payload["entries"][0]
            self.assertEqual(entry["negative_prompt"], "bad anatomy")
            self.assertEqual(entry["style"], "anime")
            self.assertEqual(entry["images"], [])

    def test_mobile_export_job_filters_selected_external_sources(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            test_db = base / "mobile-external-export.db"
            export_dir = base / "exports"
            init_db(test_db)
            now = "2026-09-03T00:00:00+00:00"
            with connect(test_db) as conn:
                conn.executemany(
                    """INSERT INTO external_sources(id,provider,title,created_at,updated_at)
                       VALUES (?, 'test', ?, ?, ?)""",
                    [("source-a", "资料库 A", now, now), ("source-b", "资料库 B", now, now)],
                )
                conn.executemany(
                    """INSERT INTO external_entries(source_id,external_id,title,prompt,updated_at)
                       VALUES (?, ?, ?, ?, ?)""",
                    [
                        ("source-a", "a-1", "A 条目", "prompt a", now),
                        ("source-b", "b-1", "B 条目", "prompt b", now),
                    ],
                )

            with patch("app.connect", partial(connect, test_db)), patch("app.MOBILE_EXPORT_DIR", export_dir):
                started = self.client.post("/api/export/mobile/jobs", json={
                    "scope": "external", "mode": "full", "source_ids": ["source-b"],
                })
                self.assertEqual(started.status_code, 202)
                job_id = started.get_json()["id"]
                deadline = time.monotonic() + 5
                status = None
                while time.monotonic() < deadline:
                    status = self.client.get(f"/api/export/mobile/jobs/{job_id}").get_json()
                    if status["status"] in {"complete", "failed"}:
                        break
                    time.sleep(0.02)
                self.assertEqual(status["status"], "complete", status)
                self.assertEqual(status["source_count"], 1)
                self.assertEqual(status["external_count"], 1)
                response = self.client.get(f"/api/export/mobile/jobs/{job_id}/download")
                self.assertEqual(response.status_code, 200)
                payload = response.get_json()
                self.assertEqual([source["id"] for source in payload["external_sources"]], ["source-b"])
                self.assertEqual([entry["source_id"] for entry in payload["external_entries"]], ["source-b"])
                self.assertIn("attachment", response.headers.get("Content-Disposition", ""))
                response.close()

    def test_backup_zip_contains_snapshot_manifest_and_originals(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            data_dir = base / "data"
            originals_dir = base / "原图"
            backup_dir = base / "backups"
            test_db = data_dir / "library.db"
            init_db(test_db)
            originals_dir.mkdir(parents=True)
            (originals_dir / "sample.png").write_bytes(b"sample-original")
            with patch("app.DATA_DIR", data_dir), patch("app.ORIGINALS_DIR", originals_dir), \
                 patch("app.BACKUP_DIR", backup_dir), patch("app.connect", partial(connect, test_db)):
                response = self.client.post("/api/backups", json={})
                self.assertEqual(response.status_code, 200)
                archive_path = Path(response.get_json()["path"])
                with zipfile.ZipFile(archive_path) as archive:
                    names = set(archive.namelist())
                self.assertIn("data/library.db", names)
                self.assertIn("manifest.json", names)
                self.assertIn("data/media/sample.png", names)

    def test_category_create_move_and_safe_delete(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "categories.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                created = self.client.post("/api/categories", json={"kind": "prompt", "path": "一级/二级"})
                self.assertEqual(created.status_code, 201)
                entry = self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "测试场景", "category": "一级/二级", "content": "test",
                })
                self.assertEqual(entry.status_code, 201)
                moved = self.client.put("/api/categories", json={
                    "kind": "prompt", "old_path": "一级", "new_path": "新一级",
                })
                self.assertEqual(moved.status_code, 200)
                entry_id = entry.get_json()["id"]
                self.assertEqual(self.client.get(f"/api/entries/{entry_id}").get_json()["category"], "新一级/二级")
                deleted = self.client.delete("/api/categories", json={"kind": "prompt", "path": "新一级"})
                self.assertEqual(deleted.status_code, 200)
                self.assertEqual(self.client.get(f"/api/entries/{entry_id}").get_json()["category"], "未分类")

    def test_category_merge_and_custom_sibling_order(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "category-drag.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                for path in ("甲", "乙", "丙"):
                    self.client.post("/api/categories", json={"kind": "prompt", "path": path})
                entry = self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "待合并", "category": "甲", "content": "test",
                }).get_json()
                merged = self.client.put("/api/categories", json={
                    "kind": "prompt", "old_path": "甲", "new_path": "乙",
                })
                self.assertEqual(merged.status_code, 200)
                self.assertEqual(self.client.get(f"/api/entries/{entry['id']}").get_json()["category"], "乙")
                reordered = self.client.put("/api/categories/reorder", json={
                    "kind": "prompt", "source_path": "丙", "target_path": "乙", "position": "before",
                })
                self.assertEqual(reordered.status_code, 200)
                navigation = self.client.get("/api/navigation").get_json()
                orders = {item["name"]: item["sort_order"] for item in navigation["categories"]["prompt"]}
                self.assertLess(orders["丙"], orders["乙"])
                self.client.post("/api/categories", json={"kind": "prompt", "path": "合成父级甲/叶子"})
                self.client.post("/api/categories", json={"kind": "prompt", "path": "合成父级乙/叶子"})
                synthesized = self.client.put("/api/categories/reorder", json={
                    "kind": "prompt", "source_path": "合成父级乙", "target_path": "合成父级甲", "position": "before",
                })
                self.assertEqual(synthesized.status_code, 200)

    def test_category_direct_place_changes_parent_and_gap_order(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "category-place.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                for path in ("待移动", "父级/子项甲", "父级/子项乙", "空父级"):
                    self.client.post("/api/categories", json={"kind": "prompt", "path": path})
                entry = self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "移动测试", "category": "待移动", "content": "test",
                }).get_json()
                placed = self.client.put("/api/categories/place", json={
                    "kind": "prompt", "source_path": "待移动", "parent_path": "父级", "before_path": "父级/子项乙",
                })
                self.assertEqual(placed.status_code, 200)
                self.assertEqual(placed.get_json()["new_path"], "父级/待移动")
                self.assertEqual(self.client.get(f"/api/entries/{entry['id']}").get_json()["category"], "父级/待移动")
                navigation = self.client.get("/api/navigation").get_json()
                orders = {item["name"]: item["sort_order"] for item in navigation["categories"]["prompt"]}
                self.assertLess(orders["父级/待移动"], orders["父级/子项乙"])
                empty_parent = self.client.put("/api/categories/place", json={
                    "kind": "prompt", "source_path": "父级/待移动", "parent_path": "空父级", "before_path": "",
                })
                self.assertEqual(empty_parent.status_code, 200)
                self.assertEqual(empty_parent.get_json()["new_path"], "空父级/待移动")

    def test_scene_upload_directory_is_flat(self):
        first = entry_media_dir("prompt", "一级/二级", "场景甲")
        second = entry_media_dir("prompt", "完全不同/四级", "场景乙")
        self.assertEqual(first, second)
        self.assertEqual(first.name, "场景")

    def test_entry_ids_honor_current_filters_for_select_all(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "select-all.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                expected = self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "目标", "category": "一级/二级", "content": "needle", "favorite": True,
                }).get_json()["id"]
                self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "未收藏", "category": "一级/二级", "content": "needle",
                })
                self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "别处", "category": "其他", "content": "needle", "favorite": True,
                })
                payload = self.client.get("/api/entries/ids", query_string={
                    "kind": "prompt", "category_prefix": "一级", "favorites_only": "1", "q": "needle",
                }).get_json()
                self.assertEqual(payload, {"ids": [expected], "total": 1})

    def test_global_search_ignores_directory_but_keeps_library_kind(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "search-scope.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                records = [
                    {"kind": "prompt", "title": "场景甲", "category": "目录甲", "content": "scope needle"},
                    {"kind": "prompt", "title": "场景乙", "category": "目录乙", "content": "scope needle"},
                    {"kind": "artist", "title": "画师甲", "category": "画师串", "content": "scope needle"},
                ]
                for record in records:
                    self.client.post("/api/entries", json=record)

                directory_result = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "q": "scope needle", "search_scope": "directory",
                    "category": "目录甲", "limit": 100,
                }).get_json()
                global_result = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "q": "scope needle", "search_scope": "all",
                    "category": "目录甲", "limit": 100,
                }).get_json()
                self.assertEqual([entry["title"] for entry in directory_result["entries"]], ["场景甲"])
                self.assertEqual({entry["title"] for entry in global_result["entries"]}, {"场景甲", "场景乙"})

    def test_entries_can_be_filtered_by_whether_they_have_images(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "image-filter.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                with_image = self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "有图场景", "category": "测试", "content": "test",
                }).get_json()["id"]
                without_image = self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "无图场景", "category": "测试", "content": "test",
                }).get_json()["id"]
                with connect(test_db) as conn:
                    conn.execute(
                        "INSERT INTO assets (path, sha256, created_at) VALUES (?, ?, ?)",
                        ("media/nai/场景/filter.png", "filter-image", "2026-08-15T00:00:00+00:00"),
                    )
                    asset_id = conn.execute("SELECT id FROM assets WHERE sha256 = 'filter-image'").fetchone()[0]
                    conn.execute("INSERT INTO entry_images (entry_id, asset_id) VALUES (?, ?)", (with_image, asset_id))

                with_result = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "image_filter": "with", "limit": 100,
                }).get_json()
                without_result = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "image_filter": "without", "limit": 100,
                }).get_json()
                temporarily_included = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "image_filter": "without",
                    "image_filter_include_ids": str(with_image), "limit": 100,
                }).get_json()
                self.assertEqual([entry["id"] for entry in with_result["entries"]], [with_image])
                self.assertEqual([entry["id"] for entry in without_result["entries"]], [without_image])
                self.assertEqual({entry["id"] for entry in temporarily_included["entries"]}, {with_image, without_image})
                navigation = self.client.get("/api/navigation").get_json()
                self.assertEqual(navigation["image_counts"]["prompt"], {
                    "all": 2, "with_images": 1, "without_images": 1,
                })

    def test_custom_groups_can_share_cards_without_changing_categories(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "groups.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                group_a = self.client.post("/api/groups", json={"kind": "prompt", "name": "测试用"}).get_json()
                group_b = self.client.post("/api/groups", json={"kind": "prompt", "name": "常用"}).get_json()
                entry = self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "场景A", "category": "正戏/测试", "content": "test",
                }).get_json()
                assigned = self.client.put(f"/api/entries/{entry['id']}/groups", json={
                    "group_ids": [group_a["id"], group_b["id"]],
                })
                self.assertEqual(assigned.status_code, 200)
                loaded = self.client.get(f"/api/entries/{entry['id']}").get_json()
                self.assertEqual(loaded["category"], "正戏/测试")
                self.assertEqual({group["name"] for group in loaded["groups"]}, {"测试用", "常用"})
                filtered = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "group_id": group_a["id"],
                }).get_json()
                self.assertEqual(filtered["total"], 1)
                self.client.delete(f"/api/groups/{group_a['id']}")
                self.assertEqual(self.client.get(f"/api/entries/{entry['id']}").status_code, 200)

    def test_folder_scan_runs_as_background_job(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            data_dir = base / "data"
            media_dir = data_dir / "media"
            scan_dir = media_dir / "待扫描"
            scan_dir.mkdir(parents=True)
            Image.new("RGB", (8, 8), "red").save(scan_dir / "sample.png")
            test_db = data_dir / "library.db"
            init_db(test_db)
            with patch("app.DATA_DIR", data_dir), patch("app.ORIGINALS_DIR", media_dir), \
                 patch("app.connect", partial(connect, test_db)):
                entry = self.client.post("/api/entries", json={
                    "kind": "prompt", "title": "后台扫描", "category": "测试", "content": "test",
                })
                entry_id = entry.get_json()["id"]
                started = self.client.post(f"/api/entries/{entry_id}/link-folder", json={"path": "待扫描"})
                self.assertEqual(started.status_code, 202)
                job_id = started.get_json()["id"]
                job = None
                for _ in range(100):
                    job = self.client.get(f"/api/jobs/{job_id}").get_json()
                    if job["status"] in ("completed", "failed"):
                        break
                    time.sleep(0.02)
                self.assertEqual(job["status"], "completed")
                self.assertEqual(job["result"]["added"], 1)

    def test_uploaded_image_is_saved_in_entry_hierarchy(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            data_dir = base / "data"
            media_dir = data_dir / "media"
            target_dir = media_dir / "nai" / "画师串" / "可改名画师"
            test_db = data_dir / "library.db"
            init_db(test_db)
            stream = BytesIO()
            Image.new("RGB", (8, 8), "blue").save(stream, "PNG")
            stream.seek(0)
            with patch("app.DATA_DIR", data_dir), patch("app.connect", partial(connect, test_db)), \
                 patch("app.entry_media_dir", return_value=target_dir):
                entry = self.client.post("/api/entries", json={
                    "kind": "artist", "title": "可改名画师", "category": "画师串", "content": "test",
                })
                entry_id = entry.get_json()["id"]
                response = self.client.post(
                    f"/api/entries/{entry_id}/images",
                    data={"images": (stream, "sample.png")}, content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.get_json()["added"], 1)
                self.assertEqual(len(list(target_dir.glob("*.png"))), 1)
                loaded = self.client.get(f"/api/entries/{entry_id}").get_json()
                self.assertEqual((loaded["images"][0]["width"], loaded["images"][0]["height"]), (8, 8))

    def test_artist_search_endpoint(self):
        response = self.client.get("/api/entries?kind=artist&limit=3")
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertLessEqual(len(payload["entries"]), 3)
        self.assertIn("total", payload)
        if payload["entries"] and payload["entries"][0]["images"]:
            image_response = self.client.get(payload["entries"][0]["images"][0]["thumbnail_url"])
            self.assertEqual(image_response.status_code, 200)
            image_response.close()
            original_response = self.client.get(payload["entries"][0]["images"][0]["url"])
            self.assertEqual(original_response.status_code, 200)
            original_response.close()

    def test_favorites_are_intersected_with_current_kind_and_category(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "favorites.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                records = [
                    {"kind": "prompt", "title": "应显示", "category": "一级/当前场景", "favorite": True},
                    {"kind": "prompt", "title": "未收藏", "category": "一级/当前场景", "favorite": False},
                    {"kind": "prompt", "title": "其他场景", "category": "一级/别处", "favorite": True},
                    {"kind": "artist", "title": "其他资料库", "category": "一级/当前场景", "favorite": True},
                ]
                for record in records:
                    response = self.client.post("/api/entries", json={**record, "content": "test"})
                    self.assertEqual(response.status_code, 201)
                response = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "category": "一级/当前场景", "favorites_only": "1",
                })
                payload = response.get_json()
                self.assertEqual(payload["total"], 1)
                self.assertEqual(payload["entries"][0]["title"], "应显示")

    def test_pinned_entries_lead_then_use_library_default_order(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "sorting.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                records = [
                    {"kind": "artist", "title": "B 画师", "rating": 10, "pinned": False},
                    {"kind": "artist", "title": "A 画师", "rating": 10, "pinned": False},
                    {"kind": "artist", "title": "低分置顶", "rating": 4, "pinned": True},
                    {"kind": "prompt", "title": "场景 B", "pinned": False},
                    {"kind": "prompt", "title": "场景 A", "pinned": False},
                ]
                for record in records:
                    response = self.client.post("/api/entries", json={**record, "category": "测试", "content": "test"})
                    self.assertEqual(response.status_code, 201)

                artists = self.client.get("/api/entries", query_string={
                    "kind": "artist", "sort": "rating_desc",
                }).get_json()["entries"]
                self.assertEqual([entry["title"] for entry in artists], ["低分置顶", "A 画师", "B 画师"])
                scenes = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "sort": "title",
                }).get_json()["entries"]
                self.assertEqual([entry["title"] for entry in scenes], ["场景 A", "场景 B"])

    def test_new_artist_is_inserted_at_end_of_its_manual_rating_band(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "manual-rating-insert.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                for title, rating in (("五星旧条目", 10), ("四星旧条目", 8), ("三星旧条目", 6)):
                    response = self.client.post("/api/entries", json={
                        "kind": "artist", "title": title, "rating": rating, "category": "测试", "content": "test",
                    })
                    self.assertEqual(response.status_code, 201)

                response = self.client.post("/api/entries", json={
                    "kind": "artist", "title": "四星新增", "rating": 8, "category": "测试", "content": "test",
                })
                self.assertEqual(response.status_code, 201)
                entries = self.client.get("/api/entries", query_string={
                    "kind": "artist", "sort": "manual", "limit": 100,
                }).get_json()["entries"]
                self.assertEqual(
                    [entry["title"] for entry in entries],
                    ["五星旧条目", "四星旧条目", "四星新增", "三星旧条目"],
                )

    def test_changed_artist_rating_moves_to_end_of_new_rating_band(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "manual-rating-update.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                ids = {}
                for title, rating in (("five-a", 10), ("five-b", 10), ("four-a", 8), ("four-b", 8)):
                    response = self.client.post("/api/entries", json={
                        "kind": "artist", "title": title, "rating": rating,
                        "category": "test", "content": "test",
                    })
                    self.assertEqual(response.status_code, 201)
                    ids[title] = response.get_json()["id"]

                # Prove an edit repositions an existing card, not only a newly created card.
                response = self.client.put(
                    f"/api/entries/{ids['four-a']}/manual-order",
                    json={"position": 1},
                )
                self.assertEqual(response.status_code, 200)
                response = self.client.put(f"/api/entries/{ids['four-a']}", json={"rating": 10})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.get_json()["manual_order"], 3)

                entries = self.client.get("/api/entries", query_string={
                    "kind": "artist", "sort": "manual", "limit": 100,
                }).get_json()["entries"]
                self.assertEqual(
                    [entry["title"] for entry in entries],
                    ["five-a", "five-b", "four-a", "four-b"],
                )

    def test_manual_order_id_moves_entry_and_renumbers_the_list(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "manual-order-id.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                ids = {}
                for title in ("scene-a", "scene-b", "scene-c"):
                    response = self.client.post("/api/entries", json={
                        "kind": "prompt", "title": title, "category": "test", "content": "test",
                    })
                    self.assertEqual(response.status_code, 201)
                    ids[title] = response.get_json()["id"]

                response = self.client.put(
                    f"/api/entries/{ids['scene-c']}/manual-order",
                    json={"position": 1},
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.get_json()["manual_order"], 1)

                entries = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "sort": "manual", "limit": 100,
                }).get_json()["entries"]
                self.assertEqual([entry["title"] for entry in entries], ["scene-c", "scene-a", "scene-b"])
                self.assertEqual([entry["manual_order"] for entry in entries], [1, 2, 3])

                response = self.client.put(
                    f"/api/entries/{ids['scene-c']}/manual-order",
                    json={"position": 999},
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.get_json()["manual_order"], 3)
                invalid = self.client.put(
                    f"/api/entries/{ids['scene-c']}/manual-order",
                    json={"position": 0},
                )
                self.assertEqual(invalid.status_code, 400)

    def test_only_pending_legacy_artists_are_migrated_into_rating_band(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "manual-rating-migration.db"
            init_db(test_db)
            now = "2026-08-20T00:00:00+00:00"
            connection = sqlite3.connect(test_db)
            try:
                connection.executemany(
                    """INSERT INTO entries
                       (kind, title, content, rating, category, manual_order, created_at, updated_at)
                       VALUES ('artist', ?, 'test', ?, '测试', ?, ?, ?)""",
                    [
                        ("五星手排", 10, 1, now, now),
                        ("四星手排", 8, 2, now, now),
                        ("三星手排", 6, 3, now, now),
                        ("四星旧版新增", 8, 0, now, now),
                    ],
                )
                connection.commit()
            finally:
                connection.close()

            # Simulate the next application start. connect() deliberately initializes each
            # database only once per process now, so a fixture inserted after that first
            # initialization must explicitly run the startup migration.
            init_db(test_db)
            with connect(test_db) as connection:
                rows = connection.execute(
                    "SELECT title, manual_order FROM entries WHERE kind = 'artist' ORDER BY manual_order"
                ).fetchall()
            self.assertEqual(
                [(row["title"], row["manual_order"]) for row in rows],
                [("五星手排", 1), ("四星手排", 2), ("四星旧版新增", 3), ("三星手排", 4)],
            )

    def test_legacy_positive_manual_orders_are_normalized_into_rating_bands_once(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "manual-rating-bands.db"
            init_db(test_db)
            now = "2026-08-20T00:00:00+00:00"
            connection = sqlite3.connect(test_db)
            try:
                connection.execute("DELETE FROM app_meta WHERE key = 'artist_manual_rating_bands'")
                connection.executemany(
                    """INSERT INTO entries
                       (kind, title, content, rating, category, manual_order, created_at, updated_at)
                       VALUES ('artist', ?, 'test', ?, 'test', ?, ?, ?)""",
                    [
                        ("unrated-a", None, 1, now, now),
                        ("five-a", 10, 2, now, now),
                        ("four-a", 8, 3, now, now),
                        ("five-b", 10, 4, now, now),
                        ("unrated-b", None, 5, now, now),
                    ],
                )
                connection.commit()
            finally:
                connection.close()

            init_db(test_db)
            with connect(test_db) as connection:
                rows = connection.execute(
                    "SELECT title, manual_order FROM entries WHERE kind = 'artist' ORDER BY manual_order"
                ).fetchall()
            self.assertEqual(
                [(row["title"], row["manual_order"]) for row in rows],
                [("five-a", 1), ("five-b", 2), ("four-a", 3), ("unrated-a", 4), ("unrated-b", 5)],
            )

    def test_chinese_title_sort_uses_pinyin_order(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "pinyin-sort.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                titles = ["两女性并排穿泳衣", "互相扣穴比赛", "亲吻嘴部特写", "侧躺阴道双头龙", "双穴假阳具床上"]
                for title in titles:
                    self.client.post("/api/entries", json={
                        "kind": "prompt", "title": title, "category": "测试", "content": "test",
                    })
                entries = self.client.get("/api/entries", query_string={
                    "kind": "prompt", "sort": "title", "limit": 100,
                }).get_json()["entries"]
                self.assertEqual(
                    [entry["title"] for entry in entries],
                    ["侧躺阴道双头龙", "互相扣穴比赛", "两女性并排穿泳衣", "亲吻嘴部特写", "双穴假阳具床上"],
                )

    def test_batch_delete_cleans_unreferenced_asset(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "batch-delete.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                ids = []
                for title in ("待删一", "待删二", "保留"):
                    response = self.client.post("/api/entries", json={
                        "kind": "prompt", "title": title, "category": "测试", "content": "test",
                    })
                    ids.append(response.get_json()["id"])
                with connect(test_db) as conn:
                    conn.execute(
                        "INSERT INTO assets (path, sha256, created_at) VALUES (?, ?, ?)",
                        ("data/media/keep.png", "keep-hash", "2026-08-15T00:00:00+00:00"),
                    )
                    asset_id = conn.execute("SELECT id FROM assets WHERE sha256 = 'keep-hash'").fetchone()[0]
                    conn.execute("INSERT INTO entry_images (entry_id, asset_id) VALUES (?, ?)", (ids[0], asset_id))

                response = self.client.post("/api/entries/batch-delete", json={"ids": ids[:2]})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.get_json()["deleted"], 2)
                self.wait_cleanup(response)
                self.assertEqual(self.client.get(f"/api/entries/{ids[0]}").status_code, 404)
                self.assertEqual(self.client.get(f"/api/entries/{ids[2]}").status_code, 200)
                with connect(test_db) as conn:
                    self.assertEqual(conn.execute("SELECT COUNT(*) FROM assets WHERE id = ?", (asset_id,)).fetchone()[0], 0)

    def test_batch_move_changes_only_selected_entries_of_the_current_kind(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "batch-move.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                prompt_ids = []
                for title in ("移动一", "移动二"):
                    response = self.client.post("/api/entries", json={
                        "kind": "prompt", "title": title, "category": "原目录", "content": "test",
                    })
                    prompt_ids.append(response.get_json()["id"])
                artist = self.client.post("/api/entries", json={
                    "kind": "artist", "title": "不应移动", "category": "画师串", "content": "test",
                }).get_json()["id"]

                response = self.client.post("/api/entries/batch-move", json={
                    "ids": [*prompt_ids, artist], "kind": "prompt", "category": "一级/二级",
                })
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.get_json()["moved"], 2)
                for entry_id in prompt_ids:
                    self.assertEqual(self.client.get(f"/api/entries/{entry_id}").get_json()["category"], "一级/二级")
                self.assertEqual(self.client.get(f"/api/entries/{artist}").get_json()["category"], "画师串")

    def test_unlink_image_only_removes_it_from_the_current_entry(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "unlink-image.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                entry_ids = []
                for title in ("当前资料", "共享资料"):
                    response = self.client.post("/api/entries", json={
                        "kind": "prompt", "title": title, "category": "测试", "content": "test",
                    })
                    entry_ids.append(response.get_json()["id"])
                with connect(test_db) as conn:
                    conn.execute(
                        "INSERT INTO assets (path, sha256, created_at) VALUES (?, ?, ?)",
                        ("media/nai/场景/shared.png", "shared-image-hash", "2026-08-15T00:00:00+00:00"),
                    )
                    asset_id = conn.execute(
                        "SELECT id FROM assets WHERE sha256 = 'shared-image-hash'"
                    ).fetchone()[0]
                    conn.executemany(
                        "INSERT INTO entry_images (entry_id, asset_id) VALUES (?, ?)",
                        [(entry_ids[0], asset_id), (entry_ids[1], asset_id)],
                    )

                response = self.client.delete(f"/api/entries/{entry_ids[0]}/images/{asset_id}")
                self.assertEqual(response.status_code, 200)
                self.wait_cleanup(response)
                self.assertEqual(self.client.get(f"/api/entries/{entry_ids[0]}").get_json()["images"], [])
                self.assertEqual(len(self.client.get(f"/api/entries/{entry_ids[1]}").get_json()["images"]), 1)
                with connect(test_db) as conn:
                    self.assertEqual(conn.execute(
                        "SELECT COUNT(*) FROM assets WHERE id = ?", (asset_id,)
                    ).fetchone()[0], 1)

    def test_cover_image_order_is_specific_to_each_entry(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "cover-image.db"
            init_db(test_db)
            with patch("app.connect", partial(connect, test_db)):
                entry_ids = []
                for title in ("更换封面", "保持顺序"):
                    response = self.client.post("/api/entries", json={
                        "kind": "prompt", "title": title, "category": "测试", "content": "test",
                    })
                    entry_ids.append(response.get_json()["id"])
                with connect(test_db) as conn:
                    conn.executemany(
                        "INSERT INTO assets (path, sha256, created_at) VALUES (?, ?, ?)",
                        [
                            ("media/nai/场景/a.png", "cover-a", "2026-08-15T00:00:00+00:00"),
                            ("media/nai/场景/b.png", "cover-b", "2026-08-15T00:00:00+00:00"),
                        ],
                    )
                    asset_ids = [row[0] for row in conn.execute("SELECT id FROM assets ORDER BY id")]
                    conn.executemany(
                        "INSERT INTO entry_images (entry_id, asset_id, sort_order) VALUES (?, ?, ?)",
                        [
                            (entry_ids[0], asset_ids[0], 0), (entry_ids[0], asset_ids[1], 1),
                            (entry_ids[1], asset_ids[0], 0), (entry_ids[1], asset_ids[1], 1),
                        ],
                    )

                response = self.client.put(
                    f"/api/entries/{entry_ids[0]}/images/cover", json={"asset_id": asset_ids[1]}
                )
                self.assertEqual(response.status_code, 200)
                first = self.client.get(f"/api/entries/{entry_ids[0]}").get_json()
                second = self.client.get(f"/api/entries/{entry_ids[1]}").get_json()
                self.assertEqual([image["id"] for image in first["images"]], [asset_ids[1], asset_ids[0]])
                self.assertEqual([image["id"] for image in second["images"]], asset_ids)


class ImageMetadataTest(unittest.TestCase):
    def make_nai_png(self, path: Path) -> None:
        info = PngImagePlugin.PngInfo()
        info.add_text("Description", "fallback prompt")
        info.add_text("Source", "NovelAI Diffusion V4.5 4BDE2A90")
        info.add_text("Comment", json.dumps({
            "prompt": "1girl, blue hair, smile",
            "uc": "bad anatomy, lowres",
            "steps": 28,
            "scale": 5.5,
            "seed": 123,
            "v4_prompt": {"caption": {
                "base_caption": "1girl, blue hair, smile",
                "char_captions": [{"char_caption": "girl, blue eyes"}],
            }},
        }))
        Image.new("RGBA", (8, 8), "blue").save(path, pnginfo=info)

    def test_novelai_png_metadata_is_split_into_prompt_fields(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "nai.png"
            self.make_nai_png(path)
            metadata = extract_image_metadata(path)
            self.assertEqual(metadata["source"], "NovelAI")
            self.assertEqual(metadata["positive_prompt"], "1girl, blue hair, smile")
            self.assertEqual(metadata["negative_prompt"], "bad anatomy, lowres")
            self.assertEqual(metadata["characters"], ["girl, blue eyes"])
            self.assertEqual(metadata["parameters"]["seed"], 123)
            self.assertEqual(metadata["parameters"]["model"], "nai-diffusion-4-5-full")

    def test_v5_source_fingerprint_distinguishes_full_and_curated(self):
        with tempfile.TemporaryDirectory() as folder:
            for source, expected in (
                ("NovelAI Diffusion V5 657484A5", "nai-diffusion-5-full"),
                ("NovelAI Diffusion V5 DE206BDA", "nai-diffusion-5-curated"),
            ):
                path = Path(folder) / (expected + ".png")
                info = PngImagePlugin.PngInfo()
                info.add_text("Source", source)
                info.add_text("Comment", json.dumps({"prompt": "test", "steps": 23}))
                Image.new("RGBA", (8, 8), "blue").save(path, pnginfo=info)
                self.assertEqual(extract_image_metadata(path)["parameters"]["model"], expected)

    def test_uploaded_metadata_is_available_from_asset_endpoint(self):
        app.config.update(TESTING=True)
        client = app.test_client()
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            data_dir = base / "data"
            media_dir = data_dir / "media"
            test_db = data_dir / "library.db"
            source = base / "source.png"
            self.make_nai_png(source)
            init_db(test_db)
            with patch("app.DATA_DIR", data_dir), patch("app.ORIGINALS_DIR", media_dir), \
                 patch("app.connect", partial(connect, test_db)), \
                 patch("app.entry_media_dir", return_value=media_dir / "nai" / "场景"):
                entry_id = client.post("/api/entries", json={
                    "kind": "prompt", "title": "元数据", "category": "测试", "content": "test",
                }).get_json()["id"]
                with source.open("rb") as stream:
                    response = client.post(
                        f"/api/entries/{entry_id}/images",
                        data={"images": (stream, "source.png")}, content_type="multipart/form-data",
                    )
                self.assertEqual(response.status_code, 200)
                asset_id = client.get(f"/api/entries/{entry_id}").get_json()["images"][0]["id"]
                metadata = client.get(f"/api/assets/{asset_id}/metadata").get_json()
                self.assertEqual(metadata["positive_prompt"], "1girl, blue hair, smile")


class LibraryAuditTest(unittest.TestCase):
    def test_audit_uses_cache_and_only_reports_images_shared_by_three_or_more(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            data_dir = base / "data"
            media_dir = data_dir / "media"
            media_dir.mkdir(parents=True)
            test_db = data_dir / "library.db"
            image_path = media_dir / "shared.png"
            info = PngImagePlugin.PngInfo()
            info.add_text("Description", "1girl, blue hair, smile")
            Image.new("RGB", (8, 8), "blue").save(image_path, pnginfo=info)
            init_db(test_db)
            now = "2026-08-17T00:00:00+00:00"
            with connect(test_db) as conn:
                conn.executemany(
                    "INSERT INTO entries (kind, title, content, category, created_at, updated_at) VALUES ('prompt', ?, ?, '测试', ?, ?)",
                    [(f"卡片{index}", "1girl, blue hair, smile", now, now) for index in range(3)],
                )
                asset_id = conn.execute(
                    "INSERT INTO assets (path, sha256, created_at) VALUES (?, ?, ?)",
                    ("media/shared.png", "audit-shared", now),
                ).lastrowid
                entry_ids = [row[0] for row in conn.execute("SELECT id FROM entries ORDER BY id")]
                conn.executemany(
                    "INSERT INTO entry_images (entry_id, asset_id) VALUES (?, ?)",
                    [(entry_id, asset_id) for entry_id in entry_ids],
                )
            _, _, first = audit_library(test_db, base / "reports")
            _, _, second = audit_library(test_db, base / "reports")
            self.assertEqual(first["summary"]["scanned"], 1)
            self.assertEqual(second["summary"]["cached"], 1)
            self.assertEqual(len(second["shared_by_three_or_more"]), 1)
            self.assertNotIn("entries_without_images", second)


class DatabaseMigrationTest(unittest.TestCase):
    def test_existing_database_gets_pinned_column_before_index(self):
        with tempfile.TemporaryDirectory() as folder:
            test_db = Path(folder) / "legacy.db"
            conn = sqlite3.connect(test_db)
            conn.execute(
                """CREATE TABLE entries (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    kind TEXT NOT NULL,
                    title TEXT NOT NULL,
                    content TEXT NOT NULL DEFAULT '',
                    negative_prompt TEXT NOT NULL DEFAULT '',
                    rating INTEGER,
                    category TEXT NOT NULL DEFAULT '未分类',
                    notes TEXT NOT NULL DEFAULT '',
                    tags TEXT NOT NULL DEFAULT '[]',
                    favorite INTEGER NOT NULL DEFAULT 0,
                    source_doc TEXT,
                    source_index INTEGER,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                )"""
            )
            conn.commit()
            conn.close()
            init_db(test_db)
            conn = sqlite3.connect(test_db)
            try:
                columns = {row[1] for row in conn.execute("PRAGMA table_info(entries)")}
                indexes = {row[1] for row in conn.execute("PRAGMA index_list(entries)")}
            finally:
                conn.close()
            self.assertIn("pinned", columns)
            self.assertIn("idx_entries_pinned", indexes)


class DocxDatabaseConverterTest(unittest.TestCase):
    def test_scene_markdown_uses_headings_as_category_hierarchy(self):
        markdown = """# 一级

## 二级

### 三级

- 场景甲：tag one, tag two
- 场景乙：tag three
"""
        self.assertEqual(
            read_markdown_document(markdown, "prompt"),
            [
                ("场景甲", "tag one, tag two", None, "一级/二级/三级"),
                ("场景乙", "tag three", None, "一级/二级/三级"),
            ],
        )

    def test_current_full_markdown_export_updates_scene_categories(self):
        markdown = """# 场景提示词库导出

导出时间：2026/8/29
资料数量：1

## 1. 场景甲

- 目录：一级/二级/三级
- 收藏：否

### 正向提示词

```text
tag one, tag two
```

### 图片

- example.png
"""
        self.assertEqual(
            read_markdown_document(markdown, "prompt"),
            [("场景甲", "tag one, tag two", None, "一级/二级/三级")],
        )

    def test_markdown_incremental_update_moves_scene_category(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            database = base / "library.db"
            init_db(database)
            now = "2026-08-29T00:00:00+00:00"
            with connect(database) as conn:
                conn.execute(
                    """INSERT INTO entries
                       (kind, title, content, category, created_at, updated_at)
                       VALUES ('prompt', '场景甲', 'same prompt', '旧分类', ?, ?)""",
                    (now, now),
                )
            records = read_markdown_document("# 新分类\n\n## 子分类\n\n- 场景甲：same prompt\n", "prompt")
            plan = build_incremental_plan(database, "prompt", records)
            self.assertEqual(plan["counts"]["modified"], 1)
            self.assertEqual(plan["modified"][0]["changes"], ["分类"])
            apply_incremental_update(database, "prompt", records, backup_dir=base / "backups")
            with connect(database) as conn:
                self.assertEqual(
                    conn.execute("SELECT category FROM entries WHERE title = '场景甲'").fetchone()[0],
                    "新分类/子分类",
                )
                paths = {
                    row[0] for row in conn.execute(
                        "SELECT path FROM categories WHERE kind = 'prompt'"
                    )
                }
            self.assertTrue({"新分类", "新分类/子分类"}.issubset(paths))

    def test_incremental_reader_resolves_colons_inside_known_titles(self):
        with tempfile.TemporaryDirectory() as folder:
            document_path = Path(folder) / "colon-title.docx"
            document = Document()
            title = "多视图：上课、吃饭、睡觉、打招呼"
            document.add_paragraph(f"{title}：{{{{multiple views}}}}, 4 views")
            document.save(document_path)
            self.assertEqual(
                read_scene_document(document_path, {title}),
                [(title, "{{multiple views}}, 4 views", None)],
            )

    def test_incremental_artist_update_preserves_images_directory_and_pinned_state(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            database = base / "data" / "library.db"
            backup_dir = base / "backups"
            init_db(database)
            now = "2026-08-18T00:00:00+00:00"
            with connect(database) as conn:
                conn.execute(
                    """
                    INSERT INTO entries
                        (kind, title, content, rating, category, favorite, pinned, tags, created_at, updated_at)
                    VALUES ('artist', '保留身份', 'old prompt', 10, '自定义目录', 1, 1, '[\"常用\"]', ?, ?)
                    """,
                    (now, now),
                )
                entry_id = conn.execute("SELECT id FROM entries WHERE title = '保留身份'").fetchone()[0]
                conn.execute(
                    "INSERT INTO assets (path, sha256, created_at) VALUES ('media/test.png', 'incremental-image', ?)",
                    (now,),
                )
                asset_id = conn.execute("SELECT id FROM assets WHERE sha256 = 'incremental-image'").fetchone()[0]
                conn.execute("INSERT INTO entry_images (entry_id, asset_id) VALUES (?, ?)", (entry_id, asset_id))
                conn.execute(
                    """
                    INSERT INTO entries (kind, title, content, category, created_at, updated_at)
                    VALUES ('artist', '通用负面', 'bad anatomy', '负面提示词', ?, ?)
                    """,
                    (now, now),
                )
                conn.execute(
                    """
                    INSERT INTO entries (kind, title, content, category, created_at, updated_at)
                    VALUES ('prompt', '不相关场景', 'scene prompt', '场景目录', ?, ?)
                    """,
                    (now, now),
                )

            records = [("保留身份", "new prompt", 9), ("新增画师", "artist prompt", None)]
            plan = build_incremental_plan(database, "artist", records)
            self.assertEqual(plan["counts"]["modified"], 1)
            self.assertEqual(plan["counts"]["added"], 1)
            applied, backup = apply_incremental_update(
                database, "artist", records, backup_dir=backup_dir,
            )
            self.assertTrue(backup.is_file())
            self.assertEqual(applied["counts"]["deleted"], 0)
            with connect(database) as conn:
                row = conn.execute(
                    "SELECT id, content, rating, category, favorite, pinned, tags FROM entries WHERE title = '保留身份'"
                ).fetchone()
                self.assertEqual(tuple(row), (entry_id, "new prompt", 9, "自定义目录", 1, 1, '["常用"]'))
                self.assertEqual(conn.execute(
                    "SELECT COUNT(*) FROM entry_images WHERE entry_id = ? AND asset_id = ?", (entry_id, asset_id)
                ).fetchone()[0], 1)
                self.assertEqual(conn.execute(
                    "SELECT COUNT(*) FROM entries WHERE title IN ('通用负面', '不相关场景')"
                ).fetchone()[0], 2)
                self.assertEqual(conn.execute(
                    "SELECT category FROM entries WHERE title = '新增画师'"
                ).fetchone()[0], "画师串")

    def test_incremental_matching_recognizes_rename_and_keeps_omitted_cards_by_default(self):
        with tempfile.TemporaryDirectory() as folder:
            database = Path(folder) / "source.db"
            init_db(database)
            now = "2026-08-18T00:00:00+00:00"
            with connect(database) as conn:
                conn.executemany(
                    """
                    INSERT INTO entries (kind, title, content, category, created_at, updated_at)
                    VALUES ('prompt', ?, ?, '原目录', ?, ?)
                    """,
                    [("旧名称", "same unique prompt", now, now), ("Word 未出现", "keep me", now, now)],
                )
            plan = build_incremental_plan(database, "prompt", [("新名称", "same unique prompt", None)])
            self.assertEqual(plan["counts"]["renamed"], 1)
            self.assertEqual(plan["counts"]["missing"], 1)
            self.assertEqual(plan["counts"]["deleted"], 0)

    def test_incremental_artist_rename_resolves_placeholder_titles_by_prompt_and_rating(self):
        with tempfile.TemporaryDirectory() as folder:
            database = Path(folder) / "artist-renames.db"
            init_db(database)
            now = "2026-08-20T00:00:00+00:00"
            with connect(database) as conn:
                conn.executemany(
                    """INSERT INTO entries
                       (kind, title, content, rating, category, created_at, updated_at)
                       VALUES ('artist', ?, ?, ?, '画师串', ?, ?)""",
                    [
                        ("未命名", "unique prompt one", 10, now, now),
                        ("未命名", "unique prompt two", 8, now, now),
                        ("占位一", "shared prompt", 9, now, now),
                        ("占位二", "shared prompt", 7, now, now),
                    ],
                )
            records = [
                ("风格一", "unique prompt one", 10),
                ("风格二", "unique prompt two", 8),
                ("光泽CG·01", "shared prompt", 9),
                ("光泽CG·02", "shared prompt", 7),
            ]
            plan = build_incremental_plan(database, "artist", records)
            self.assertTrue(plan["can_apply"])
            self.assertEqual(plan["counts"], {
                "added": 0, "modified": 0, "renamed": 4, "unchanged": 0,
                "missing": 0, "deleted": 0, "ambiguous": 0,
            })

    def test_incremental_update_page_previews_then_applies_same_document(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            data_dir = base / "data"
            backup_dir = base / "backups"
            database = data_dir / "library.db"
            init_db(database)
            now = "2026-08-18T00:00:00+00:00"
            with connect(database) as conn:
                conn.execute(
                    """
                    INSERT INTO entries (kind, title, content, category, created_at, updated_at)
                    VALUES ('prompt', '网页更新', 'old prompt', '保留目录', ?, ?)
                    """,
                    (now, now),
                )
            document = Document()
            document.add_paragraph("网页更新：new prompt")
            stream = BytesIO()
            document.save(stream)
            payload = stream.getvalue()
            app.config.update(TESTING=True)
            client = app.test_client()
            with patch("app.DATA_DIR", data_dir), patch("app.BACKUP_DIR", backup_dir):
                page = client.get("/converter")
                self.assertEqual(page.status_code, 200)
                page.close()
                preview = client.post(
                    "/api/converter/preview",
                    data={"kind": "prompt", "delete_missing": "0", "document": (BytesIO(payload), "场景.docx")},
                    content_type="multipart/form-data",
                )
                self.assertEqual(preview.status_code, 200)
                preview_data = preview.get_json()
                self.assertEqual(preview_data["plan"]["counts"]["modified"], 1)
                applied = client.post(
                    "/api/converter/apply",
                    data={
                        "kind": "prompt", "delete_missing": "0",
                        "document_digest": preview_data["document_digest"],
                        "plan_token": preview_data["plan_token"],
                        "document": (BytesIO(payload), "场景.docx"),
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(applied.status_code, 200)
            with connect(database) as conn:
                row = conn.execute("SELECT content, category FROM entries WHERE title = '网页更新'").fetchone()
                self.assertEqual(tuple(row), ("new prompt", "保留目录"))

    def test_import_preview_separates_added_modified_removed_and_ambiguous(self):
        with tempfile.TemporaryDirectory() as folder:
            database = Path(folder) / "source.db"
            init_db(database)
            now = "2026-08-17T00:00:00+00:00"
            with connect(database) as conn:
                conn.executemany(
                    "INSERT INTO entries (kind, title, content, category, created_at, updated_at) VALUES ('prompt', ?, ?, '测试', ?, ?)",
                    [
                        ("不变", "same prompt", now, now),
                        ("修改", "old prompt", now, now),
                        ("删除", "removed prompt", now, now),
                    ],
                )
            preview = preview_import(
                database,
                [("不变", "same prompt", None), ("修改", "new prompt", None), ("新增", "added prompt", None)],
                [],
            )
            self.assertEqual(preview["totals"], {
                "added": 1, "modified": 1, "removed": 1, "unchanged": 1, "ambiguous": 0,
            })

    def test_database_and_separate_word_documents_round_trip_without_images(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            source_db = base / "source.db"
            backup_dir = base / "backups"
            init_db(source_db)
            now = "2026-08-15T00:00:00+00:00"
            with connect(source_db) as conn:
                conn.executemany(
                    """
                    INSERT INTO entries (kind, title, content, rating, category, source_doc, source_index, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        ("prompt", "场景甲", "tag one, tag two", None, "一级/二级/三级", "test", 1, now, now),
                        ("artist", "画师甲", "artist one", 9, "画师串", "test", 2, now, now),
                        ("artist", "未评分画师", "artist two", None, "画师串", "test", 3, now, now),
                        ("artist", "通用负面", "bad anatomy", None, "负面提示词", "test", 4, now, now),
                    ],
                )

            scenes_docx, artists_docx = export_database(source_db, backup_dir, "2026-08-15")
            self.assertEqual(scenes_docx.name, "word - 2026-08-15 - v1 - 场景.docx")
            self.assertEqual(artists_docx.name, "word - 2026-08-15 - v1 - 画师串.docx")
            scene_document = Document(scenes_docx)
            nonempty = [paragraph for paragraph in scene_document.paragraphs if paragraph.text.strip()]
            self.assertEqual(
                [(paragraph.text, paragraph.style.style_id) for paragraph in nonempty[:3]],
                [("一级", "Heading1"), ("二级", "Heading2"), ("三级", "Heading3")],
            )
            self.assertAlmostEqual(nonempty[1].paragraph_format.left_indent.cm, 0.74, places=1)
            self.assertAlmostEqual(nonempty[2].paragraph_format.left_indent.cm, 1.48, places=1)
            self.assertAlmostEqual(nonempty[3].paragraph_format.left_indent.cm, 2.22, places=1)
            rebuilt_db = import_documents(scenes_docx, artists_docx, backup_dir, "2026-08-15")
            self.assertEqual(rebuilt_db.name, "数据库 - 2026-08-15 - v1.db")
            self.assertTrue((backup_dir / "数据库 - 2026-08-15 - v1 - 差异预览.json").is_file())
            with connect(rebuilt_db) as conn:
                rows = conn.execute(
                    "SELECT kind, title, content, rating, category FROM entries ORDER BY kind DESC, title"
                ).fetchall()
                records = [tuple(row) for row in rows]
                self.assertIn(("prompt", "场景甲", "tag one, tag two", None, "一级/二级/三级"), records)
                self.assertIn(("artist", "画师甲", "artist one", 9, "画师串"), records)
                self.assertIn(("artist", "未评分画师", "artist two", None, "画师串"), records)
                self.assertNotIn("通用负面", [row[1] for row in records])
                self.assertEqual(conn.execute("SELECT COUNT(*) FROM entry_images").fetchone()[0], 0)

            second_scenes, _ = export_database(source_db, backup_dir, "2026-08-15")
            self.assertIn("v2", second_scenes.name)


if __name__ == "__main__":
    unittest.main()
