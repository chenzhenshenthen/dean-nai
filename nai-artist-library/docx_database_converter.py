from __future__ import annotations

import argparse
import json
import re
import sqlite3
import unicodedata
from datetime import date
from pathlib import Path
from typing import BinaryIO

from docx import Document
from docx.shared import Cm

from database import DB_PATH, ROOT, init_db, utc_now
from importer import normalize_text, rating_from_text, split_mapping
from library_media import cleanup_detached


BACKUP_DIR = ROOT / "backups"


def next_version(backup_dir: Path, prefix: str, day: str) -> int:
    backup_dir.mkdir(parents=True, exist_ok=True)
    pattern = re.compile(rf"^{re.escape(prefix)} - {re.escape(day)} - v(\d+)(?:\D.*)?$")
    versions = []
    for path in backup_dir.iterdir():
        match = pattern.match(path.stem)
        if match:
            versions.append(int(match.group(1)))
    return max(versions, default=0) + 1


def rating_label(rating: int | None) -> str:
    if rating is None:
        return "未评分"
    return f"{rating / 2:g}★"


def parse_rating_heading(value: str) -> tuple[bool, int | None]:
    text = value.strip()
    if text in {"未评分", "无评分", "尚未评分"}:
        return True, None
    numeric = re.fullmatch(r"([0-5](?:\.5)?)\s*(?:★|星)?", text)
    if numeric:
        score = float(numeric.group(1))
        rating = int(score * 2)
        return (1 <= rating <= 10), rating if 1 <= rating <= 10 else None
    if re.fullmatch(r"[★☆]{1,5}", text):
        return True, rating_from_text(text)
    return False, None


def read_entries(database_path: Path, kind: str) -> list[sqlite3.Row]:
    database_path = database_path.resolve()
    if not database_path.is_file():
        raise FileNotFoundError(f"数据库不存在：{database_path}")
    connection = sqlite3.connect(f"{database_path.as_uri()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    try:
        if kind == "prompt":
            return connection.execute(
                """SELECT title, content, rating, category FROM entries
                   WHERE kind = 'prompt'
                   ORDER BY category COLLATE NOCASE, title COLLATE NOCASE, id"""
            ).fetchall()
        return connection.execute(
            """
            SELECT title, content, rating, category FROM entries
            WHERE kind = 'artist' AND category <> '负面提示词'
            ORDER BY rating IS NULL, rating DESC, title COLLATE NOCASE, id
            """
        ).fetchall()
    finally:
        connection.close()


def write_scene_document(rows: list[sqlite3.Row], target: Path) -> None:
    document = Document()
    emitted: set[str] = set()
    for row in rows:
        components = [part.strip() for part in str(row["category"] or "\u672a\u5206\u7c7b").split("/") if part.strip()]
        if not components:
            components = ["\u672a\u5206\u7c7b"]
        for depth, component in enumerate(components, start=1):
            path = "/".join(components[:depth])
            if path in emitted:
                continue
            heading = document.add_heading(component, level=min(depth, 9))
            heading.paragraph_format.left_indent = Cm(0.74 * (depth - 1))
            emitted.add(path)
        paragraph = document.add_paragraph(f"{row['title']}：{row['content']}")
        paragraph.paragraph_format.left_indent = Cm(0.74 * len(components))
    document.save(target)


def write_artist_document(rows: list[sqlite3.Row], target: Path) -> None:
    document = Document()
    grouped: dict[int | None, list[sqlite3.Row]] = {}
    for row in rows:
        grouped.setdefault(row["rating"], []).append(row)
    ratings = sorted((rating for rating in grouped if rating is not None), reverse=True)
    if None in grouped:
        ratings.append(None)
    for group_index, rating in enumerate(ratings):
        heading = document.add_paragraph()
        heading.add_run(rating_label(rating)).bold = True
        for row in grouped[rating]:
            paragraph = document.add_paragraph(f"{row['title']}：{row['content']}")
            paragraph.paragraph_format.left_indent = Cm(0.74)
        if group_index < len(ratings) - 1:
            document.add_paragraph()
    document.save(target)


def export_database(database_path: Path, backup_dir: Path = BACKUP_DIR, day: str | None = None) -> tuple[Path, Path]:
    day = day or date.today().isoformat()
    version = next_version(backup_dir, "word", day)
    base = f"word - {day} - v{version}"
    scene_target = backup_dir / f"{base} - 场景.docx"
    artist_target = backup_dir / f"{base} - 画师串.docx"
    write_scene_document(read_entries(database_path, "prompt"), scene_target)
    write_artist_document(read_entries(database_path, "artist"), artist_target)
    return scene_target, artist_target


DocumentSource = Path | str | BinaryIO
IncrementalRecord = tuple[str, str, int | None] | tuple[str, str, int | None, str]


def record_parts(record: IncrementalRecord) -> tuple[str, str, int | None, str | None]:
    title, content, rating = record[:3]
    raw_category = record[3] if len(record) > 3 else None
    category = str(raw_category).strip() if raw_category is not None and str(raw_category).strip() else None
    return title, content, rating, category


def split_document_mapping(text: str, known_titles: set[str] | None = None) -> tuple[str, str] | None:
    if known_titles:
        # Exported titles may themselves contain ':' or '：'. Prefer an exact
        # current database title over guessing which colon is the delimiter.
        for title in sorted(known_titles, key=len, reverse=True):
            for delimiter in ("：", ":"):
                prefix = f"{title}{delimiter}"
                if text.startswith(prefix):
                    return normalize_text(title), normalize_text(text[len(prefix):])
    return split_mapping(text)


def _docx_heading_level(paragraph) -> int:
    style_id = str(getattr(paragraph.style, "style_id", "") or "")
    match = re.search(r"Heading(\d+)", style_id, re.IGNORECASE)
    if match:
        return max(1, int(match.group(1)))
    indent = paragraph.paragraph_format.left_indent
    if indent is not None:
        return max(1, round(indent.cm / 0.74) + 1)
    return 1


def read_scene_document(path: DocumentSource, known_titles: set[str] | None = None) -> list[IncrementalRecord]:
    records: list[IncrementalRecord] = []
    category_stack: list[str] = []
    for paragraph in Document(path).paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        mapping = split_document_mapping(text, known_titles)
        if not mapping:
            level = _docx_heading_level(paragraph)
            if level <= len(category_stack):
                category_stack = category_stack[:level - 1]
            category_stack.append(normalize_text(text))
            continue
        title, content = mapping
        if category_stack:
            records.append((title, content, None, "/".join(category_stack)))
        else:
            records.append((title, content, None))
    return records


def _markdown_text(source: Path | str | bytes | BinaryIO) -> str:
    if isinstance(source, Path):
        return source.read_text(encoding="utf-8-sig")
    if isinstance(source, bytes):
        return source.decode("utf-8-sig")
    if hasattr(source, "read"):
        value = source.read()
        return value.decode("utf-8-sig") if isinstance(value, bytes) else str(value)
    path = Path(source)
    return path.read_text(encoding="utf-8-sig") if path.is_file() else source


def _read_full_export_markdown(text: str, kind: str) -> list[IncrementalRecord] | None:
    if not re.search(r"(?m)^##\s+\d+\.\s+", text) or "- 目录：" not in text:
        return None
    records: list[IncrementalRecord] = []
    current: dict | None = None
    section = ""
    in_fence = False
    fence = chr(96) * 3

    def finish() -> None:
        nonlocal current
        if not current or not current["content"]:
            current = None
            return
        category = current["category"] or ("画师串" if kind == "artist" else "未分类")
        if kind == "artist" and category == "负面提示词":
            current = None
            return
        content = normalize_text("\n".join(current["content"]))
        if kind == "prompt":
            records.append((current["title"], content, None, category))
        else:
            records.append((current["title"], content, current["rating"]))
        current = None

    for raw_line in text.splitlines():
        stripped = raw_line.strip()
        entry_heading = re.fullmatch(r"##\s+\d+\.\s+(.+?)\s*", stripped)
        if entry_heading:
            finish()
            current = {
                "title": normalize_text(entry_heading.group(1)),
                "category": None, "rating": None, "content": [],
            }
            section = ""
            in_fence = False
            continue
        if current is None:
            continue
        if stripped.startswith("- 目录："):
            current["category"] = normalize_text(stripped.removeprefix("- 目录："))
            continue
        if stripped.startswith("- 评分："):
            value = stripped.removeprefix("- 评分：").strip()
            match = re.match(r"(\d+(?:\.\d+)?)\s*/\s*5", value)
            current["rating"] = round(float(match.group(1)) * 2) if match else None
            continue
        if stripped == "### 正向提示词":
            section = "positive"
            continue
        if stripped.startswith("### "):
            section = ""
            continue
        if stripped.startswith(fence):
            in_fence = not in_fence
            continue
        if in_fence and section == "positive":
            current["content"].append(raw_line)
    finish()
    if not records:
        raise ValueError("Markdown 完整导出中没有识别到可更新的资料")
    return records


def read_markdown_document(
    source: Path | str | bytes | BinaryIO,
    kind: str,
    known_titles: set[str] | None = None,
) -> list[IncrementalRecord]:
    if kind not in {"prompt", "artist"}:
        raise ValueError("\u66f4\u65b0\u7c7b\u578b\u5fc5\u987b\u662f prompt \u6216 artist")
    text_content = _markdown_text(source)
    full_export = _read_full_export_markdown(text_content, kind)
    if full_export is not None:
        return full_export
    records: list[IncrementalRecord] = []
    category_stack: list[str] = []
    current_rating: int | None = None
    rating_heading_seen = False
    in_fence = False
    for line_number, raw_line in enumerate(text_content.splitlines(), start=1):
        text = raw_line.strip()
        if text.startswith("\x60\x60\x60") or text.startswith("~~~"):
            in_fence = not in_fence
            continue
        if in_fence or not text or text.startswith("<!--"):
            continue
        heading = re.fullmatch(r"(#{1,9})\s+(.+?)\s*#*", text)
        if heading:
            level = len(heading.group(1))
            label = heading.group(2).strip().strip("*_\x60")
            if kind == "artist":
                is_rating, parsed_rating = parse_rating_heading(label)
                if not is_rating:
                    raise ValueError(f"\u753b\u5e08\u4e32 Markdown \u7b2c {line_number} \u884c\u7684\u6807\u9898\u4e0d\u662f\u8bc4\u5206\uff1a{label[:80]}")
                current_rating = parsed_rating
                rating_heading_seen = True
            else:
                if level <= len(category_stack):
                    category_stack = category_stack[:level - 1]
                category_stack.append(normalize_text(label))
            continue
        candidate = re.sub(r"^(?:[-+*]|\d+[.)])\s+", "", text)
        mapping = split_document_mapping(candidate, known_titles)
        if not mapping:
            raise ValueError(f"Markdown \u7b2c {line_number} \u884c\u65e0\u6cd5\u8bc6\u522b\uff1a{text[:80]}")
        title, content = mapping
        if kind == "prompt" and category_stack:
            records.append((title, content, None, "/".join(category_stack)))
        else:
            records.append((title, content, current_rating if rating_heading_seen else None))
    if not records:
        raise ValueError("Markdown \u4e2d\u6ca1\u6709\u8bc6\u522b\u5230\u4efb\u4f55\u8d44\u6599")
    return records


def read_incremental_document(
    source: DocumentSource | bytes,
    kind: str,
    suffix: str,
    known_titles: set[str] | None = None,
) -> list[IncrementalRecord]:
    if suffix.lower() in {".md", ".markdown"}:
        return read_markdown_document(source, kind, known_titles)
    if suffix.lower() != ".docx":
        raise ValueError("\u53ea\u652f\u6301 .docx\u3001.md \u6216 .markdown \u6587\u4ef6")
    return read_scene_document(source, known_titles) if kind == "prompt" else read_artist_document(source, known_titles)


def read_artist_document(path: DocumentSource, known_titles: set[str] | None = None) -> list[tuple[str, str, int | None]]:
    records = []
    current_rating: int | None = None
    rating_heading_seen = False
    for paragraph in Document(path).paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        is_heading, parsed_rating = parse_rating_heading(text)
        if is_heading:
            current_rating = parsed_rating
            rating_heading_seen = True
            continue
        if text in {"待识别风格："}:
            # Optional organizational label used by manually curated artist
            # documents; it is not an artist entry and does not change rating.
            continue
        mapping = split_document_mapping(text, known_titles)
        if not mapping:
            raise ValueError(f"画师串 Word 中存在无法识别的段落：{text[:80]}")
        title, content = mapping
        records.append((title, content, current_rating if rating_heading_seen else None))
    return records


def insert_records(database_path: Path, scenes: list[IncrementalRecord], artists: list[IncrementalRecord]) -> None:
    init_db(database_path)
    now = utc_now()
    connection = sqlite3.connect(database_path)
    try:
        with connection:
            connection.executemany(
                """
                INSERT INTO entries
                    (kind, title, content, rating, category, source_doc, source_index, created_at, updated_at)
                VALUES ('prompt', ?, ?, ?, ?, '转换-场景.docx', ?, ?, ?)
                """,
                [
                    (title, content, rating, category or "未分类", index, now, now)
                    for index, record in enumerate(scenes)
                    for title, content, rating, category in [record_parts(record)]
                ],
            )
            connection.executemany(
                """
                INSERT INTO entries
                    (kind, title, content, rating, category, source_doc, source_index, created_at, updated_at)
                VALUES ('artist', ?, ?, ?, '画师串', '转换-画师串.docx', ?, ?, ?)
                """,
                [(title, content, rating, index, now, now) for index, (title, content, rating) in enumerate(artists)],
            )
    finally:
        connection.close()


def _database_records(database_path: Path, kind: str) -> list[IncrementalRecord]:
    if not database_path.is_file():
        return []
    connection = sqlite3.connect(database_path)
    try:
        columns = "title, content, rating, category" if kind == "prompt" else "title, content, rating"
        return [tuple(row) for row in connection.execute(
            f"SELECT {columns} FROM entries WHERE kind = ? AND NOT (kind = 'artist' AND category = '负面提示词')",
            (kind,),
        ).fetchall()]
    finally:
        connection.close()


def _compare_records(
    current: list[IncrementalRecord],
    incoming: list[IncrementalRecord],
) -> dict:
    current_by_title: dict[str, list[tuple[str, str, int | None]]] = {}
    incoming_by_title: dict[str, list[tuple[str, str, int | None]]] = {}
    for record in current:
        current_by_title.setdefault(record[0].strip(), []).append(record)
    for record in incoming:
        incoming_by_title.setdefault(record[0].strip(), []).append(record)
    result = {"added": [], "modified": [], "removed": [], "unchanged": [], "ambiguous": []}
    for title in sorted(set(current_by_title) | set(incoming_by_title)):
        before = current_by_title.get(title, [])
        after = incoming_by_title.get(title, [])
        # 只有 Word 真正涉及这个标题时，数据库中的重名才会造成匹配歧义。
        # 仅存在于数据库、且默认保留的历史重名不应阻止其他资料更新。
        if len(after) > 1 or (after and len(before) > 1):
            result["ambiguous"].append({
                "title": title, "current_count": len(before), "incoming_count": len(after),
            })
        elif not before:
            result["added"].append({"title": title})
        elif not after:
            result["removed"].append({"title": title})
        else:
            before_parts = record_parts(before[0])
            after_parts = record_parts(after[0])
            effective_after = (
                after_parts[0], after_parts[1], after_parts[2],
                after_parts[3] if after_parts[3] is not None else before_parts[3],
            )
            if before_parts == effective_after:
                result["unchanged"].append({"title": title})
                continue
            changed = []
            if before[0][1] != after[0][1]:
                changed.append("提示词")
            if before[0][2] != after[0][2]:
                changed.append("评分")
            if before_parts[3] != effective_after[3]:
                changed.append("分类")
            result["modified"].append({"title": title, "changed": changed})
    result["counts"] = {key: len(value) for key, value in result.items() if isinstance(value, list)}
    return result


def preview_import(
    database_path: Path,
    scenes: list[IncrementalRecord],
    artists: list[IncrementalRecord],
) -> dict:
    scene_diff = _compare_records(_database_records(database_path, "prompt"), scenes)
    artist_diff = _compare_records(_database_records(database_path, "artist"), artists)
    totals = {
        key: scene_diff["counts"][key] + artist_diff["counts"][key]
        for key in ("added", "modified", "removed", "unchanged", "ambiguous")
    }
    return {
        "format": "nai-library-import-preview",
        "compared_database": str(database_path.resolve()),
        "scenes": scene_diff,
        "artists": artist_diff,
        "totals": totals,
    }


def preview_summary(preview: dict) -> str:
    totals = preview["totals"]
    return (
        f"新增 {totals['added']} · 修改 {totals['modified']} · 删除/未导入 {totals['removed']} · "
        f"不变 {totals['unchanged']} · 重名待复核 {totals['ambiguous']}"
    )


def _normalize_title(value: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", value).strip()).casefold()


def _normalize_content(value: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", value).strip()).casefold()


def _incremental_existing(database_path: Path, kind: str) -> list[dict]:
    if kind not in {"prompt", "artist"}:
        raise ValueError("资料库类型必须是 prompt 或 artist")
    connection = sqlite3.connect(database_path)
    connection.row_factory = sqlite3.Row
    try:
        where = "kind = ?"
        if kind == "artist":
            where += " AND category <> '负面提示词'"
        return [dict(row) for row in connection.execute(
            f"SELECT id, title, content, rating, category FROM entries WHERE {where} ORDER BY id", (kind,)
        ).fetchall()]
    finally:
        connection.close()


def incremental_known_titles(database_path: Path, kind: str) -> set[str]:
    return {item["title"] for item in _incremental_existing(database_path, kind)}


def build_incremental_plan(
    database_path: Path,
    kind: str,
    records: list[IncrementalRecord],
    *,
    delete_missing: bool = False,
) -> dict:
    """按标题优先、唯一内容其次匹配 Word 与现有资料，生成不改数据库的更新计划。"""
    existing = _incremental_existing(database_path, kind)
    incoming = []
    for index, record in enumerate(records):
        title, content, rating, category = record_parts(record)
        incoming.append({
            "index": index, "title": title.strip(), "content": content.strip(),
            "rating": rating, "category": category,
        })
    plan = {
        "format": "nai-library-incremental-preview",
        "kind": kind,
        "delete_missing": delete_missing,
        "added": [], "modified": [], "renamed": [], "unchanged": [],
        "missing": [], "deleted": [], "ambiguous": [],
    }

    existing_titles: dict[str, list[dict]] = {}
    incoming_titles: dict[str, list[dict]] = {}
    for item in existing:
        existing_titles.setdefault(_normalize_title(item["title"]), []).append(item)
    for item in incoming:
        incoming_titles.setdefault(_normalize_title(item["title"]), []).append(item)

    matched_existing: set[int] = set()
    matched_incoming: set[int] = set()
    ambiguous_existing: set[int] = set()
    ambiguous_incoming: set[int] = set()

    for key in sorted(set(existing_titles) | set(incoming_titles)):
        before = existing_titles.get(key, [])
        after = incoming_titles.get(key, [])
        if len(before) == 1 and len(after) == 1:
            old, new = before[0], after[0]
            matched_existing.add(old["id"])
            matched_incoming.add(new["index"])
            changes = []
            if _normalize_content(old["content"]) != _normalize_content(new["content"]):
                changes.append("提示词")
            if kind == "artist" and old["rating"] != new["rating"]:
                changes.append("评分")
            new_category = new["category"] if kind == "prompt" and new["category"] is not None else old["category"]
            if kind == "prompt" and old["category"] != new_category:
                changes.append("分类")
            target = plan["modified"] if changes else plan["unchanged"]
            target.append({
                "id": old["id"], "title": old["title"], "new_title": new["title"],
                "old_content": old["content"], "new_content": new["content"],
                "old_rating": old["rating"], "new_rating": new["rating"], "changes": changes,
                "old_category": old["category"], "new_category": new_category,
            })

    remaining_existing = [
        item for item in existing
        if item["id"] not in matched_existing and item["id"] not in ambiguous_existing
    ]
    remaining_incoming = [
        item for item in incoming
        if item["index"] not in matched_incoming and item["index"] not in ambiguous_incoming
    ]
    existing_content: dict[str, list[dict]] = {}
    incoming_content: dict[str, list[dict]] = {}
    for item in remaining_existing:
        key = _normalize_content(item["content"])
        if key:
            existing_content.setdefault(key, []).append(item)
    for item in remaining_incoming:
        key = _normalize_content(item["content"])
        if key:
            incoming_content.setdefault(key, []).append(item)

    def match_by_content(old: dict, new: dict) -> None:
        matched_existing.add(old["id"])
        matched_incoming.add(new["index"])
        changes = []
        if old["title"] != new["title"]:
            changes.append("名称")
        if kind == "artist" and old["rating"] != new["rating"]:
            changes.append("评分")
        new_category = new["category"] if kind == "prompt" and new["category"] is not None else old["category"]
        if kind == "prompt" and old["category"] != new_category:
            changes.append("分类")
        item = {
            "id": old["id"], "title": old["title"], "new_title": new["title"],
            "old_content": old["content"], "new_content": new["content"],
            "old_rating": old["rating"], "new_rating": new["rating"], "changes": changes,
            "old_category": old["category"], "new_category": new_category,
        }
        if "名称" in changes:
            plan["renamed"].append(item)
        elif changes:
            plan["modified"].append(item)
        else:
            plan["unchanged"].append(item)

    for key in sorted(set(existing_content) & set(incoming_content)):
        before = [item for item in existing_content[key] if item["id"] not in matched_existing]
        after = [item for item in incoming_content[key] if item["index"] not in matched_incoming]
        if len(before) == 1 and len(after) == 1:
            match_by_content(before[0], after[0])
            continue
        if kind == "artist":
            # Two artist strings can intentionally share the same prompt. A
            # distinct rating provides a stable identity for rename documents.
            for rating in sorted({item["rating"] for item in before} | {item["rating"] for item in after}, key=lambda value: (value is None, value)):
                old_at_rating = [item for item in before if item["rating"] == rating and item["id"] not in matched_existing]
                new_at_rating = [item for item in after if item["rating"] == rating and item["index"] not in matched_incoming]
                if len(old_at_rating) == 1 and len(new_at_rating) == 1:
                    match_by_content(old_at_rating[0], new_at_rating[0])

    # Only unresolved duplicate titles/content on both sides are genuinely
    # ambiguous. Duplicate old placeholder names alone must not block unique
    # prompt-based renames.
    remaining_existing = [item for item in existing if item["id"] not in matched_existing]
    remaining_incoming = [item for item in incoming if item["index"] not in matched_incoming]
    remaining_existing_titles: dict[str, list[dict]] = {}
    remaining_incoming_titles: dict[str, list[dict]] = {}
    for item in remaining_existing:
        remaining_existing_titles.setdefault(_normalize_title(item["title"]), []).append(item)
    for item in remaining_incoming:
        remaining_incoming_titles.setdefault(_normalize_title(item["title"]), []).append(item)
    for key in sorted(set(remaining_existing_titles) & set(remaining_incoming_titles)):
        before, after = remaining_existing_titles[key], remaining_incoming_titles[key]
        plan["ambiguous"].append({
            "reason": "标题重复且无法按提示词唯一匹配", "title": (after or before)[0]["title"],
            "database_ids": [item["id"] for item in before],
            "word_indexes": [item["index"] for item in after],
        })
        ambiguous_existing.update(item["id"] for item in before)
        ambiguous_incoming.update(item["index"] for item in after)

    remaining_existing_content: dict[str, list[dict]] = {}
    remaining_incoming_content: dict[str, list[dict]] = {}
    for item in remaining_existing:
        if item["id"] not in ambiguous_existing:
            remaining_existing_content.setdefault(_normalize_content(item["content"]), []).append(item)
    for item in remaining_incoming:
        if item["index"] not in ambiguous_incoming:
            remaining_incoming_content.setdefault(_normalize_content(item["content"]), []).append(item)
    for key in sorted((set(remaining_existing_content) & set(remaining_incoming_content)) - {""}):
        before, after = remaining_existing_content[key], remaining_incoming_content[key]
        if before and after:
            plan["ambiguous"].append({
                "reason": "提示词及评分重复", "title": after[0]["title"],
                "database_ids": [item["id"] for item in before],
                "word_indexes": [item["index"] for item in after],
            })
            ambiguous_existing.update(item["id"] for item in before)
            ambiguous_incoming.update(item["index"] for item in after)

    for item in incoming:
        if item["index"] in matched_incoming or item["index"] in ambiguous_incoming:
            continue
        plan["added"].append({
            "word_index": item["index"], "title": item["title"], "new_content": item["content"],
            "new_rating": item["rating"], "new_category": item["category"],
        })
    for item in existing:
        if item["id"] in matched_existing or item["id"] in ambiguous_existing:
            continue
        target = "deleted" if delete_missing else "missing"
        plan[target].append({
            "id": item["id"], "title": item["title"], "old_content": item["content"],
            "old_rating": item["rating"], "category": item["category"],
        })

    plan["counts"] = {
        key: len(plan[key])
        for key in ("added", "modified", "renamed", "unchanged", "missing", "deleted", "ambiguous")
    }
    plan["can_apply"] = not plan["ambiguous"]
    return plan


def _backup_database(database_path: Path, backup_dir: Path) -> Path:
    day = date.today().isoformat()
    version = next_version(backup_dir, "数据库", day)
    target = backup_dir / f"数据库 - {day} - v{version} - 增量前.db"
    source = sqlite3.connect(database_path)
    backup = sqlite3.connect(target)
    try:
        source.backup(backup)
    finally:
        backup.close()
        source.close()
    return target


def apply_incremental_update(
    database_path: Path,
    kind: str,
    records: list[IncrementalRecord],
    *,
    delete_missing: bool = False,
    backup_dir: Path = BACKUP_DIR,
) -> tuple[dict, Path]:
    plan = build_incremental_plan(database_path, kind, records, delete_missing=delete_missing)
    if plan["ambiguous"]:
        raise ValueError("存在重复标题或其他歧义，请先修正 Word；数据库没有发生变化")
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = _backup_database(database_path, backup_dir)
    connection = sqlite3.connect(database_path)
    connection.row_factory = sqlite3.Row
    connection.execute('PRAGMA foreign_keys = ON')
    try:
        now = utc_now()
        with connection:
            for item in [*plan["modified"], *plan["renamed"]]:
                if kind == "artist":
                    connection.execute(
                        "UPDATE entries SET title = ?, content = ?, rating = ?, updated_at = ? WHERE id = ? AND kind = ?",
                        (item["new_title"], item["new_content"], item["new_rating"], now, item["id"], kind),
                    )
                else:
                    connection.execute(
                        "UPDATE entries SET title = ?, content = ?, category = ?, updated_at = ? WHERE id = ? AND kind = ?",
                        (item["new_title"], item["new_content"], item["new_category"], now, item["id"], kind),
                    )
            default_category = "画师串" if kind == "artist" else "未分类"
            for item in plan["added"]:
                category = item["new_category"] or default_category
                connection.execute(
                    """
                    INSERT INTO entries
                        (kind, title, content, rating, category, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (kind, item["title"], item["new_content"], item["new_rating"] if kind == "artist" else None,
                     category, now, now),
                )
            if kind == "prompt":
                categories = {
                    item.get("new_category")
                    for item in [*plan["added"], *plan["modified"], *plan["renamed"]]
                    if item.get("new_category")
                }
                for category in categories:
                    parts = [part.strip() for part in category.split("/") if part.strip()]
                    for depth in range(1, len(parts) + 1):
                        connection.execute(
                            "INSERT OR IGNORE INTO categories (kind, path, created_at) VALUES (?, ?, ?)",
                            (kind, "/".join(parts[:depth]), now),
                        )
            if delete_missing and plan["deleted"]:
                removed_ids = [item['id'] for item in plan['deleted']]
                placeholders = ','.join('?' for _ in removed_ids)
                asset_ids = [row[0] for row in connection.execute(f'SELECT DISTINCT asset_id FROM entry_images WHERE entry_id IN ({placeholders})', removed_ids)]
                connection.executemany(
                    "DELETE FROM entries WHERE id = ? AND kind = ?",
                    [(item["id"], kind) for item in plan["deleted"]],
                )
                plan['media_cleanup'] = cleanup_detached(connection, asset_ids)
    finally:
        connection.close()
    return plan, backup_path


def import_documents(
    scene_path: Path,
    artist_path: Path,
    backup_dir: Path = BACKUP_DIR,
    day: str | None = None,
    compare_database: Path = DB_PATH,
) -> Path:
    if not scene_path.is_file():
        raise FileNotFoundError(f"场景 Word 不存在：{scene_path}")
    if not artist_path.is_file():
        raise FileNotFoundError(f"画师串 Word 不存在：{artist_path}")
    day = day or date.today().isoformat()
    version = next_version(backup_dir, "数据库", day)
    target = backup_dir / f"数据库 - {day} - v{version}.db"
    scenes = read_scene_document(scene_path)
    artists = read_artist_document(artist_path)
    preview = preview_import(compare_database, scenes, artists)
    preview.update({"created_at": utc_now(), "target_database": str(target.resolve())})
    preview_path = backup_dir / f"数据库 - {day} - v{version} - 差异预览.json"
    preview_path.write_text(json.dumps(preview, ensure_ascii=False, indent=2), encoding="utf-8")
    insert_records(target, scenes, artists)
    return target


def interactive() -> int:
    print("\nDOCX / 数据库本地转换工具")
    print("1. 数据库转换为场景 Word + 画师串 Word")
    print("2. 场景 Word + 画师串 Word 转换为新数据库")
    print("3. 用场景 Word 增量更新当前数据库")
    print("4. 用画师串 Word 增量更新当前数据库")
    choice = input("请选择 1、2、3 或 4：").strip()
    if choice == "1":
        raw = input(f"数据库路径（直接回车使用 {DB_PATH}）：").strip()
        scene, artist = export_database(Path(raw) if raw else DB_PATH)
        print(f"场景 Word：{scene}\n画师串 Word：{artist}")
        return 0
    if choice == "2":
        scene = Path(input("场景 Word 路径：").strip().strip('"'))
        artist = Path(input("画师串 Word 路径：").strip().strip('"'))
        preview = preview_import(DB_PATH, read_scene_document(scene), read_artist_document(artist))
        print(f"导入差异预览：{preview_summary(preview)}")
        if input("确认生成新的备份数据库吗？输入 y 继续：").strip().casefold() not in {"y", "yes"}:
            print("已取消，未生成数据库。")
            return 0
        target = import_documents(scene, artist)
        print(f"新数据库：{target}\n差异报告：{target.with_name(target.stem + ' - 差异预览.json')}")
        return 0
    if choice in {"3", "4"}:
        kind = "prompt" if choice == "3" else "artist"
        label = "场景" if kind == "prompt" else "画师串"
        document_path = Path(input(f"{label} Word 路径：").strip().strip('"'))
        known_titles = incremental_known_titles(DB_PATH, kind)
        records = read_incremental_document(document_path, kind, document_path.suffix, known_titles)
        plan = build_incremental_plan(DB_PATH, kind, records)
        print(json.dumps(plan["counts"], ensure_ascii=False, indent=2))
        if not plan["can_apply"]:
            print("存在匹配歧义，未修改数据库。")
            return 1
        if input("确认按以上预览增量更新吗？输入 y 继续：").strip().casefold() not in {"y", "yes"}:
            print("已取消，数据库没有变化。")
            return 0
        _, backup_path = apply_incremental_update(DB_PATH, kind, records)
        print(f"更新完成；更新前数据库备份：{backup_path}")
        return 0
    print("选择无效。")
    return 2


def main() -> int:
    parser = argparse.ArgumentParser(description="在本地将 NAI 资料库与两个简洁 DOCX 文件相互转换。")
    subparsers = parser.add_subparsers(dest="command")
    export_parser = subparsers.add_parser("export", help="数据库转换为 Word")
    export_parser.add_argument("--database", type=Path, default=DB_PATH)
    import_parser = subparsers.add_parser("import", help="Word 转换为新数据库")
    import_parser.add_argument("--scenes", type=Path, required=True)
    import_parser.add_argument("--artists", type=Path, required=True)
    import_parser.add_argument("--compare-database", type=Path, default=DB_PATH, help="差异预览所比较的现有数据库")
    update_parser = subparsers.add_parser("update", help="用一个 Word 文件增量更新场景或画师串")
    update_parser.add_argument("--kind", choices=("prompt", "artist"), required=True)
    update_parser.add_argument("--document", type=Path, required=True)
    update_parser.add_argument("--database", type=Path, default=DB_PATH)
    update_parser.add_argument("--delete-missing", action="store_true", help="删除 Word 中未出现的同类卡片")
    update_parser.add_argument("--preview-only", action="store_true", help="只输出预览，不修改数据库")
    args = parser.parse_args()
    try:
        if args.command == "export":
            scene, artist = export_database(args.database)
            print(f"场景 Word：{scene}\n画师串 Word：{artist}")
            return 0
        if args.command == "import":
            scenes = read_scene_document(args.scenes)
            artists = read_artist_document(args.artists)
            print(f"导入差异预览：{preview_summary(preview_import(args.compare_database, scenes, artists))}")
            print(f"新数据库：{import_documents(args.scenes, args.artists, compare_database=args.compare_database)}")
            return 0
        if args.command == "update":
            known_titles = incremental_known_titles(args.database, args.kind)
            records = read_incremental_document(args.document, args.kind, args.document.suffix, known_titles)
            plan = build_incremental_plan(
                args.database, args.kind, records, delete_missing=args.delete_missing,
            )
            print(json.dumps(plan, ensure_ascii=False, indent=2))
            if args.preview_only:
                return 0 if plan["can_apply"] else 1
            if not plan["can_apply"]:
                print("存在匹配歧义，数据库没有变化。")
                return 1
            _, backup_path = apply_incremental_update(
                args.database, args.kind, records, delete_missing=args.delete_missing,
            )
            print(f"更新完成；更新前数据库备份：{backup_path}")
            return 0
        return interactive()
    except Exception as error:
        print(f"转换失败：{error}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
